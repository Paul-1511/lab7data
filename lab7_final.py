from __future__ import annotations

import argparse
import gc
import hashlib
import json
import math
import os
import platform
import re
import shutil
import subprocess
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime, timezone
from importlib import metadata
from pathlib import Path
from typing import Any, Callable

ROOT = Path(os.environ.get("LAB7_ROOT", Path(__file__).resolve().parent))
DATA_DIR = Path(os.environ.get("LAB7_DATA_DIR", ROOT / "datos"))
DICT_DIR = DATA_DIR / "diccionarios"
PROCESSED_DIR = Path(os.environ.get("LAB7_PROCESSED_DIR", ROOT / "data_processed"))
OUTPUT_DIR = Path(os.environ.get("LAB7_OUTPUT_DIR", ROOT / "outputs"))
FOUNDATION_OUT = OUTPUT_DIR / "foundation"
PERIOD_AUDIT_DIR = FOUNDATION_OUT / "periods"
FOUNDATION_DATA = PROCESSED_DIR / "foundation"
SELECTED_DIR = FOUNDATION_DATA / "selected_typed"
ELIGIBLE_PERIOD_DIR = FOUNDATION_DATA / "eligible_by_period"
ELIGIBLE_2025 = PROCESSED_DIR / "eligible_2025.parquet"
ELIGIBLE_2026 = PROCESSED_DIR / "eligible_2026.parquet"
PREFLIGHT_PARQUET = PROCESSED_DIR / "_preflight_parquet"
MODELS_DIR = Path(os.environ.get("LAB7_MODELS_DIR", ROOT / "models"))
EDA_OUT = OUTPUT_DIR / "eda"
EDA_FIGURES = EDA_OUT / "figures"
CLUSTERS_2025 = PROCESSED_DIR / "clusters_2025.parquet"
KMEANS_MODEL = MODELS_DIR / "kmeans_2025"
MODELS_VALIDATION_OUT = OUTPUT_DIR / "models"
VALIDATION_PREDICTIONS = PROCESSED_DIR / "predictions_validation_2025T4.parquet"
LR_VALIDATION_MODEL = MODELS_DIR / "validation_linear_regression"
RF_VALIDATION_MODEL = MODELS_DIR / "validation_random_forest"

REQUIRED_COLUMNS: tuple[str, ...] = (
    "ANIO",
    "TRIMESTRE",
    "DOMINIO",
    "NUM_HOGAR",
    "NUM_PERSONA",
    "FACTOR",
    "P02A03",
    "P03A03A",
    "P05C07A",
    "P05C07B",
    "P05H01A",
    "P05C16",
    "OCUPADOS",
    "P05D01",
)
COMPLETE_COLUMNS = ("ANIO", "TRIMESTRE", "DOMINIO", "NUM_HOGAR", "NUM_PERSONA", "FACTOR", "P02A03")
CONDITIONAL_LABOR_COLUMNS = ("P05C07A", "P05C07B", "P05H01A", "P05C16", "OCUPADOS")
INTEGER_CODE_COLUMNS = ("ANIO", "TRIMESTRE", "DOMINIO", "P03A03A", "P05C16", "OCUPADOS")
LONG_CODE_COLUMNS = ("NUM_HOGAR", "NUM_PERSONA")
DOUBLE_COLUMNS = ("FACTOR", "P02A03", "P05C07A", "P05C07B", "P05H01A", "P05D01")
AUDITED_CODES: dict[str, tuple[int, ...]] = {
    "DOMINIO": (1, 2, 3),
    "P03A03A": tuple(range(8)),
    "P05C16": tuple(range(1, 10)),
    "OCUPADOS": (1,),
}
ELIGIBLE_CODES: dict[str, tuple[int, ...]] = {
    "nivel_educativo": tuple(range(8)),
    "categoria_ocupacional": (1, 2, 3, 4),
    "dominio": (1, 2, 3),
}
UNKNOWN_LABEL = "DESCONOCIDO"
CATEGORY_LABELS: dict[str, dict[str, str]] = {
    "nivel_educativo": {
        "0": "Ninguno",
        "1": "Preprimaria",
        "2": "Primaria",
        "3": "Básico",
        "4": "Diversificado",
        "5": "Superior",
        "6": "Maestría",
        "7": "Doctorado",
    },
    "categoria_ocupacional": {
        "1": "Empleado de gobierno",
        "2": "Empleado de empresa privada",
        "3": "Jornalero o peón",
        "4": "Servicio doméstico",
    },
    "dominio": {"1": "Urbano Metropolitano", "2": "Resto Urbano", "3": "Rural Nacional"},
}
CATEGORY_TITLES = {
    "nivel_educativo": "Nivel educativo",
    "categoria_ocupacional": "Categoría ocupacional",
    "dominio": "Dominio",
}
NUMERIC_VARIABLES = ("salario_mensual", "edad", "antiguedad", "horas_semanales")
NUMERIC_TITLES = {
    "salario_mensual": "Salario mensual (Q)",
    "edad": "Edad (años)",
    "antiguedad": "Antigüedad (años)",
    "horas_semanales": "Horas semanales",
}
CLUSTER_FEATURES = ("edad", "antiguedad", "horas_semanales", "nivel_educativo_ordinal")
PERCENTILES = (0.25, 0.5, 0.75, 0.95)
SEED = 3066
SAMPLE_SIZE = 5000
K_RANGE = (2, 3, 4, 5)
KMEANS_MAX_ITER = 100
MODEL_COLUMNS = (
    "salario_mensual",
    "edad",
    "antiguedad",
    "horas_semanales",
    "nivel_educativo",
    "categoria_ocupacional",
    "dominio",
)
ELIGIBLE_SCHEMA: dict[str, str] = {
    "id_registro": "string",
    "archivo_origen": "string",
    "periodo_archivo": "string",
    "anio_archivo": "int",
    "trimestre_calendario": "int",
    "fila_origen": "int",
    "ANIO": "int",
    "TRIMESTRE": "int",
    "NUM_HOGAR": "bigint",
    "NUM_PERSONA": "bigint",
    "FACTOR": "double",
    "edad": "double",
    "antiguedad": "double",
    "horas_semanales": "double",
    "salario_mensual": "double",
    "nivel_educativo": "string",
    "categoria_ocupacional": "string",
    "dominio": "string",
    "P03A03A": "int",
    "P05C16": "int",
    "codigo_dominio": "int",
    "OCUPADOS": "int",
}
CONTROL_TOTALS = {
    "desarrollo_2025T1_T3": (("2025T1", "2025T2", "2025T3"), 40_361),
    "validacion_2025T4": (("2025T4",), 12_664),
    "reentrenamiento_2025": (("2025T1", "2025T2", "2025T3", "2025T4"), 53_025),
    "prueba_2026T1": (("2026T1",), 13_258),
}
TRAIN_PERIODS, TRAIN_ROWS = CONTROL_TOTALS["desarrollo_2025T1_T3"]
VALIDATION_PERIOD = CONTROL_TOTALS["validacion_2025T4"][0][0]
VALIDATION_ROWS = CONTROL_TOTALS["validacion_2025T4"][1]
TARGET = "salario_mensual"
NUMERIC_PREDICTORS = ("edad", "antiguedad", "horas_semanales")
CATEGORICAL_PREDICTORS = ("nivel_educativo", "categoria_ocupacional", "dominio")
PREDICTORS = (*NUMERIC_PREDICTORS, *CATEGORICAL_PREDICTORS)
LR_MAX_ITER = 200
LR_REG_PARAMS = (0.0, 10.0, 100.0, 1000.0)
LR_CONFIGS: list[dict[str, Any]] = [
    {"config_id": f"lr_reg_{reg}", "regParam": reg, "elasticNetParam": 0.0} for reg in LR_REG_PARAMS
]
RF_MAX_BINS = 32
RF_CONFIGS: list[dict[str, Any]] = [
    {"config_id": f"rf_t{trees}_d{depth}", "numTrees": trees, "maxDepth": depth} for trees, depth in ((50, 6), (100, 8), (100, 10))
]
RAW_TOTAL_CONTROLS = {"2025": 203_676, "2026": 49_843, "total": 253_519}
SENTINELS = {"", "nan", "none", "null", "na", "n/a", "#n/a", ".", "-", "inf", "+inf", "-inf", "infinity", "-infinity"}
INTEGRAL_TEXT = re.compile(r"^([+-]?\d+)\.0+$")
THOUSANDS_TEXT = re.compile(r"^[+-]?\d{1,3}(,\d{3})+(\.\d+)?$")
MIB = 1024 * 1024
MIN_RECOMMENDED_MEMORY_GIB = 3.5


def expected_missing(salary: int, education: int, labor: int) -> dict[str, int]:
    missing = {column: 0 for column in COMPLETE_COLUMNS}
    missing["P05D01"] = salary
    missing["P03A03A"] = education
    missing.update({column: labor for column in CONDITIONAL_LABOR_COLUMNS})
    return missing


@dataclass(frozen=True)
class SourceSpec:
    periodo: str
    archivo: str
    hoja: str
    anio: int
    trimestre: int
    filas: int
    columnas: int
    tamano_mib: float
    elegibles: int
    trimestre_original: dict[int, int]
    faltantes: dict[str, int]
    diccionario: str

    @property
    def path(self) -> Path:
        return DATA_DIR / self.archivo

    @property
    def dictionary_path(self) -> Path:
        return DICT_DIR / self.diccionario


SOURCES: tuple[SourceSpec, ...] = (
    SourceSpec(
        "2025T1", "Personas_2025T1.xlsx", "Personas_ENEIC_T1_2025", 2025, 1, 51_588, 270, 45.5, 13_419,
        {2: 51_588}, expected_missing(38_169, 6_928, 29_315), "Diccionario_Personas_ENEIC_I-2025.xlsx",
    ),
    SourceSpec(
        "2025T2", "Personas_2025T2.xlsx", "Personas ENEIC T2-2025", 2025, 2, 51_167, 270, 16.1, 13_492,
        {3: 50_992, 2: 175}, expected_missing(37_675, 6_678, 28_824), "Diccionario_Personas_ENEIC_II-2025.xlsx",
    ),
    SourceSpec(
        "2025T3", "Personas_2025T3.xlsx", "Personas ENEIC T3-2025", 2025, 3, 51_583, 270, 45.4, 13_450,
        {4: 51_583}, expected_missing(38_133, 6_548, 29_210), "Diccionario-Personas-ENEIC-III-2025.xlsx",
    ),
    SourceSpec(
        "2025T4", "Personas_2025T4.xlsx", "Base de datos Personas ENEIC IV", 2025, 4, 49_338, 302, 49.0, 12_664,
        {5: 49_338}, expected_missing(36_674, 6_190, 28_105), "Diccionario-Personas-ENEIC-IV-2025.xlsx",
    ),
    SourceSpec(
        "2026T1", "Personas_2026T1.xlsx", "Personas_ENEIC_T1-2026", 2026, 1, 49_843, 270, 43.9, 13_258,
        {6: 49_843}, expected_missing(36_585, 6_011, 28_109), "Diccionario-Personas-ENEIC-I-2026.xlsx",
    ),
)
SOURCES_BY_PERIOD = {source.periodo: source for source in SOURCES}


class FoundationError(RuntimeError):
    pass


@dataclass
class CheckLog:
    failures: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    passed: int = 0

    def require(self, condition: bool, message: str) -> None:
        if condition:
            self.passed += 1
        else:
            self.failures.append(message)
            print(f"  FALLA: {message}")

    def warn(self, condition: bool, message: str) -> None:
        if condition:
            self.passed += 1
        else:
            self.warnings.append(message)
            print(f"  ADVERTENCIA: {message}")

    @property
    def ok(self) -> bool:
        return not self.failures

    def as_dict(self) -> dict[str, Any]:
        return {"ok": self.ok, "controles_aprobados": self.passed, "fallas": self.failures, "advertencias": self.warnings}


@dataclass
class WorkbookExtract:
    hoja: str
    hojas: list[str]
    hojas_visibles: list[str]
    columnas_encabezado: int
    encabezados_vacios: int
    encabezados_duplicados: list[str]
    filas_datos: int
    filas_vacias_omitidas: int
    rows: list[tuple[Any, ...]]


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def relative(path: Path) -> str:
    try:
        return path.resolve().relative_to(ROOT.resolve()).as_posix()
    except ValueError:
        return path.as_posix()


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=False), encoding="utf-8")
    temporary.replace(path)


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def sha256_file(path: Path, chunk_size: int = 4 * MIB) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(chunk_size), b""):
            digest.update(chunk)
    return digest.hexdigest()


def package_version(name: str) -> str | None:
    try:
        return metadata.version(name)
    except metadata.PackageNotFoundError:
        return None


def java_version() -> str | None:
    executable = shutil.which("java")
    if executable is None:
        return None
    completed = subprocess.run([executable, "-version"], capture_output=True, text=True, check=False)
    lines = (completed.stderr or completed.stdout).strip().splitlines()
    return lines[0] if lines else None


def memory_info() -> dict[str, float | None]:
    info: dict[str, float | None] = {"mem_total_gib": None, "mem_available_gib": None, "cgroup_limit_gib": None}
    meminfo = Path("/proc/meminfo")
    if meminfo.exists():
        for line in meminfo.read_text().splitlines():
            key, _, value = line.partition(":")
            if key in ("MemTotal", "MemAvailable"):
                gib = round(int(value.split()[0]) / (1024 * 1024), 2)
                info["mem_total_gib" if key == "MemTotal" else "mem_available_gib"] = gib
    cgroup = Path("/sys/fs/cgroup/memory.max")
    if cgroup.exists():
        raw = cgroup.read_text().strip()
        if raw.isdigit():
            info["cgroup_limit_gib"] = round(int(raw) / (1024**3), 2)
    return info


def environment_info() -> dict[str, Any]:
    return {
        "python": sys.version.split()[0],
        "plataforma": platform.platform(),
        "cpu": os.cpu_count(),
        "java": java_version(),
        "paquetes": {
            name: package_version(name)
            for name in ("pyspark", "pandas", "pyarrow", "numpy", "openpyxl", "matplotlib", "seaborn")
        },
        "ps_disponible": shutil.which("ps") is not None,
        "memoria": memory_info(),
        "spark_master": os.environ.get("LAB7_SPARK_MASTER", "local[2]"),
        "spark_driver_memory": os.environ.get("LAB7_DRIVER_MEMORY", "1g"),
    }


def normalize_header(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def normalize_cell(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return None
        return str(int(value)) if value.is_integer() else repr(value)
    text = str(value).strip()
    if text.lower() in SENTINELS:
        return None
    if THOUSANDS_TEXT.match(text):
        text = text.replace(",", "")
    integral = INTEGRAL_TEXT.match(text)
    return integral.group(1) if integral else text


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def read_workbook(path: Path, columns: tuple[str, ...] = REQUIRED_COLUMNS) -> WorkbookExtract:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = list(workbook.sheetnames)
        visible = [sheet.title for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
        if len(visible) != 1:
            raise FoundationError(f"{path.name}: se esperaba una hoja visible y se encontraron {visible}")
        rows_iterator = workbook[visible[0]].iter_rows(values_only=True)
        header_row = next(rows_iterator, None)
        if header_row is None:
            raise FoundationError(f"{path.name}: hoja sin encabezado")
        header = [normalize_header(value) for value in header_row]
        while header and header[-1] is None:
            header.pop()
        names = [name for name in header if name is not None]
        duplicates = sorted({name for name in names if names.count(name) > 1})
        missing = [column for column in columns if column not in header]
        if missing:
            raise FoundationError(f"{path.name}: faltan columnas requeridas {missing}")
        if any(column in duplicates for column in columns):
            raise FoundationError(f"{path.name}: columnas requeridas duplicadas {duplicates}")
        indexes = [header.index(column) for column in columns]
        width = len(header)
        rows: list[tuple[Any, ...]] = []
        skipped = 0
        excel_row = 1
        for raw in rows_iterator:
            excel_row += 1
            if raw is None or all(is_blank(value) for value in raw[:width]):
                skipped += 1
                continue
            rows.append(
                tuple(normalize_cell(raw[index]) if index < len(raw) else None for index in indexes) + (excel_row,)
            )
        return WorkbookExtract(
            hoja=visible[0],
            hojas=sheets,
            hojas_visibles=visible,
            columnas_encabezado=width,
            encabezados_vacios=len(header) - len(names),
            encabezados_duplicados=duplicates,
            filas_datos=len(rows),
            filas_vacias_omitidas=skipped,
            rows=rows,
        )
    finally:
        workbook.close()


def build_spark(app_name: str):
    from pyspark.sql import SparkSession

    return (
        SparkSession.builder.appName(app_name)
        .master(os.environ.get("LAB7_SPARK_MASTER", "local[2]"))
        .config("spark.driver.memory", os.environ.get("LAB7_DRIVER_MEMORY", "1g"))
        .config("spark.sql.shuffle.partitions", "4")
        .config("spark.default.parallelism", "2")
        .config("spark.sql.session.timeZone", "UTC")
        .config("spark.sql.ansi.enabled", "false")
        .config("spark.ui.enabled", "false")
        .config("spark.sql.parquet.compression.codec", "snappy")
        .config("spark.sql.execution.arrow.pyspark.enabled", "false")
        .getOrCreate()
    )


def raw_schema():
    from pyspark.sql.types import IntegerType, StringType, StructField, StructType

    fields = [StructField(column, StringType(), True) for column in REQUIRED_COLUMNS]
    return StructType(fields + [StructField("fila_origen", IntegerType(), False)])


def finite_double(column: str):
    from pyspark.sql import functions as F

    value = F.col(column).cast("double")
    return F.when(value.isNotNull() & ~F.isnan(value) & (F.abs(value) < F.lit(float("inf"))), value)


def integral_code(column: str, data_type: str):
    from pyspark.sql import functions as F

    value = finite_double(column)
    return F.when(value == F.floor(value), value.cast(data_type))


def typed_expression(column: str):
    if column in LONG_CODE_COLUMNS:
        return integral_code(column, "bigint")
    if column in INTEGER_CODE_COLUMNS:
        return integral_code(column, "int")
    return finite_double(column)


def with_provenance(frame, source: SourceSpec):
    from pyspark.sql import functions as F

    return frame.select(
        F.lit(source.archivo).alias("archivo_origen"),
        F.lit(source.periodo).alias("periodo_archivo"),
        F.lit(source.anio).cast("int").alias("anio_archivo"),
        F.lit(source.trimestre).cast("int").alias("trimestre_calendario"),
        "fila_origen",
        *REQUIRED_COLUMNS,
    )


def typed_frame(raw_frame):
    from pyspark.sql import functions as F

    provenance = ("archivo_origen", "periodo_archivo", "anio_archivo", "trimestre_calendario")
    return raw_frame.select(
        *provenance,
        F.col("fila_origen").cast("int").alias("fila_origen"),
        *[typed_expression(column).alias(column) for column in REQUIRED_COLUMNS],
    )


def raw_quality(raw_frame) -> tuple[dict[str, int], dict[str, int]]:
    from pyspark.sql import functions as F

    aggregations = []
    for column in REQUIRED_COLUMNS:
        aggregations.append(F.sum(F.when(F.col(column).isNull(), 1).otherwise(0)).alias(f"nulo__{column}"))
        invalid = F.col(column).isNotNull() & typed_expression(column).isNull()
        aggregations.append(F.sum(F.when(invalid, 1).otherwise(0)).alias(f"invalido__{column}"))
    row = raw_frame.agg(*aggregations).first().asDict()
    missing = {column: int(row[f"nulo__{column}"] or 0) for column in REQUIRED_COLUMNS}
    invalid = {column: int(row[f"invalido__{column}"] or 0) for column in REQUIRED_COLUMNS}
    return missing, invalid


def value_counts(frame, column: str) -> dict[str, int]:
    rows = frame.groupBy(column).count().collect()
    counts = {("NULO" if row[column] is None else str(row[column])): int(row["count"]) for row in rows}
    return dict(sorted(counts.items(), key=lambda item: (item[0] == "NULO", item[0])))


def key_audit(frame) -> dict[str, int]:
    from pyspark.sql import functions as F

    total = frame.count()
    null_keys = frame.filter(F.col("NUM_HOGAR").isNull() | F.col("NUM_PERSONA").isNull()).count()
    distinct_keys = frame.select("NUM_HOGAR", "NUM_PERSONA").distinct().count()
    distinct_rows = frame.select(*REQUIRED_COLUMNS).distinct().count()
    return {
        "filas": total,
        "llaves_nulas": null_keys,
        "llaves_distintas": distinct_keys,
        "llaves_duplicadas": total - distinct_keys,
        "duplicados_exactos_columnas_seleccionadas": total - distinct_rows,
    }


def tenure_expression():
    from pyspark.sql import functions as F

    return F.col("P05C07A") + F.col("P05C07B") / F.lit(12.0)


def filter_steps() -> list[tuple[str, str, Callable[[], Any]]]:
    from pyspark.sql import functions as F

    def age_valid():
        return F.col("P02A03").isNotNull() & (F.col("P02A03") >= 15)

    def employed_wage_earner():
        return (F.col("OCUPADOS") == 1) & F.col("P05C16").isin(1, 2, 3, 4)

    def positive_salary():
        return F.col("P05D01").isNotNull() & (F.col("P05D01") > 0)

    def tenure_components():
        months = F.col("P05C07B")
        return (
            F.col("P05C07A").isNotNull()
            & (F.col("P05C07A") >= 0)
            & months.isNotNull()
            & (months == F.floor(months))
            & (months >= 0)
            & (months <= 11)
        )

    def tenure_within_age():
        return tenure_expression() <= F.col("P02A03")

    def weekly_hours():
        return F.col("P05H01A").isNotNull() & (F.col("P05H01A") > 0) & (F.col("P05H01A") <= 168)

    return [
        ("1_edad", "P02A03 numerico, finito y >= 15", age_valid),
        ("2_ocupado_categoria", "OCUPADOS == 1 y P05C16 en {1,2,3,4}", employed_wage_earner),
        ("3_salario", "P05D01 numerico, finito y > 0", positive_salary),
        ("4_antiguedad_componentes", "P05C07A >= 0 y P05C07B entero entre 0 y 11", tenure_components),
        ("5_antiguedad_edad", "P05C07A + P05C07B/12 <= edad", tenure_within_age),
        ("6_horas", "P05H01A numerico, finito, > 0 y <= 168", weekly_hours),
    ]


def apply_filters(frame) -> tuple[Any, list[dict[str, Any]]]:
    from pyspark.sql import functions as F

    current = frame
    before = current.count()
    steps = []
    for key, description, condition in filter_steps():
        current = current.filter(F.coalesce(condition(), F.lit(False)))
        after = current.count()
        steps.append({"paso": key, "regla": description, "antes": before, "excluidos": before - after, "despues": after})
        before = after
    return current, steps


def code_label(column: str, codes: tuple[int, ...]):
    from pyspark.sql import functions as F

    return F.when(F.col(column).isin(*codes), F.col(column).cast("string")).otherwise(F.lit(UNKNOWN_LABEL))


def eligible_frame(filtered):
    from pyspark.sql import functions as F

    return filtered.select(
        F.concat_ws(
            "_", F.col("periodo_archivo"), F.col("NUM_HOGAR").cast("string"), F.col("NUM_PERSONA").cast("string")
        ).alias("id_registro"),
        "archivo_origen",
        "periodo_archivo",
        "anio_archivo",
        "trimestre_calendario",
        "fila_origen",
        "ANIO",
        "TRIMESTRE",
        "NUM_HOGAR",
        "NUM_PERSONA",
        "FACTOR",
        F.col("P02A03").alias("edad"),
        tenure_expression().alias("antiguedad"),
        F.col("P05H01A").alias("horas_semanales"),
        F.col("P05D01").alias("salario_mensual"),
        code_label("P03A03A", ELIGIBLE_CODES["nivel_educativo"]).alias("nivel_educativo"),
        code_label("P05C16", ELIGIBLE_CODES["categoria_ocupacional"]).alias("categoria_ocupacional"),
        code_label("DOMINIO", ELIGIBLE_CODES["dominio"]).alias("dominio"),
        "P03A03A",
        "P05C16",
        F.col("DOMINIO").alias("codigo_dominio"),
        "OCUPADOS",
    )


def unknown_counts(frame) -> dict[str, int]:
    from pyspark.sql import functions as F

    row = frame.agg(
        *[F.sum(F.when(F.col(column) == UNKNOWN_LABEL, 1).otherwise(0)).alias(column) for column in ELIGIBLE_CODES]
    ).first()
    return {column: int(row[column] or 0) for column in ELIGIBLE_CODES}


def unrecognized_codes(frame) -> dict[str, int]:
    from pyspark.sql import functions as F

    row = frame.agg(
        *[
            F.sum(F.when(F.col(column).isNotNull() & ~F.col(column).isin(*codes), 1).otherwise(0)).alias(column)
            for column, codes in AUDITED_CODES.items()
        ]
    ).first()
    return {column: int(row[column] or 0) for column in AUDITED_CODES}


def period_checks(audit: dict[str, Any], source: SourceSpec, checks: CheckLog) -> None:
    label = source.periodo
    checks.require(audit["hoja"] == source.hoja, f"{label}: hoja {audit['hoja']!r} distinta de {source.hoja!r}")
    checks.require(len(audit["hojas_visibles"]) == 1, f"{label}: numero de hojas visibles distinto de 1")
    checks.require(audit["filas_datos"] == source.filas, f"{label}: filas {audit['filas_datos']} != {source.filas}")
    checks.require(
        audit["columnas_encabezado"] == source.columnas,
        f"{label}: columnas {audit['columnas_encabezado']} != {source.columnas}",
    )
    checks.require(audit["encabezados_vacios"] == 0, f"{label}: encabezados vacios")
    checks.require(not audit["encabezados_duplicados"], f"{label}: encabezados duplicados")
    keys = audit["llaves"]
    checks.require(keys["llaves_nulas"] == 0, f"{label}: llaves nulas {keys['llaves_nulas']}")
    checks.require(keys["llaves_duplicadas"] == 0, f"{label}: llaves duplicadas {keys['llaves_duplicadas']}")
    checks.require(
        keys["duplicados_exactos_columnas_seleccionadas"] == 0,
        f"{label}: duplicados exactos {keys['duplicados_exactos_columnas_seleccionadas']}",
    )
    expected_quarter = {str(key): value for key, value in source.trimestre_original.items()}
    checks.require(
        audit["trimestre_original"] == expected_quarter,
        f"{label}: TRIMESTRE original {audit['trimestre_original']} != {expected_quarter}",
    )
    checks.warn(
        audit["anio_original_valores"] == {str(source.anio): source.filas},
        f"{label}: ANIO original {audit['anio_original_valores']}",
    )
    for column, expected in source.faltantes.items():
        observed = audit["faltantes_crudos"][column]
        checks.warn(observed == expected, f"{label}: faltantes crudos {column} {observed} != control {expected}")
    for column, count in audit["valores_no_convertibles"].items():
        checks.warn(count == 0, f"{label}: {count} valores no numericos en {column}")
    for column, count in audit["codigos_no_reconocidos"].items():
        checks.warn(count == 0, f"{label}: {count} codigos fuera de rango en {column}")
    steps = audit["exclusiones"]
    checks.require(steps[0]["antes"] == audit["filas_datos"], f"{label}: el primer filtro no parte de todas las filas")
    for step in steps:
        checks.require(step["antes"] - step["excluidos"] == step["despues"], f"{label}: aritmetica {step['paso']}")
    for previous, following in zip(steps, steps[1:]):
        checks.require(previous["despues"] == following["antes"], f"{label}: cadena {previous['paso']}")
    checks.require(steps[-1]["despues"] == audit["elegibles"], f"{label}: ultimo filtro distinto de elegibles")
    checks.require(
        audit["elegibles"] == source.elegibles, f"{label}: elegibles {audit['elegibles']} != control {source.elegibles}"
    )
    for column, count in audit["desconocidos_elegibles"].items():
        checks.require(count == 0, f"{label}: {count} {UNKNOWN_LABEL} en {column}")


def parquet_complete(path: Path) -> bool:
    return (path / "_SUCCESS").exists()


def period_is_current(source: SourceSpec, digest: str) -> bool:
    audit_path = PERIOD_AUDIT_DIR / f"{source.periodo}.json"
    if not audit_path.exists():
        return False
    audit = read_json(audit_path)
    return (
        audit.get("sha256") == digest
        and audit.get("columnas_requeridas") == list(REQUIRED_COLUMNS)
        and parquet_complete(SELECTED_DIR / source.periodo)
        and parquet_complete(ELIGIBLE_PERIOD_DIR / source.periodo)
    )


def process_source(spark, source: SourceSpec, digest: str) -> dict[str, Any]:
    started = time.perf_counter()
    print(f"[{source.periodo}] leyendo {source.archivo} (solo {len(REQUIRED_COLUMNS)} columnas)")
    extract = read_workbook(source.path)
    rows = extract.rows
    extract.rows = []
    print(f"[{source.periodo}] filas={extract.filas_datos} columnas={extract.columnas_encabezado} hoja={extract.hoja!r}")
    raw_frame = with_provenance(spark.createDataFrame(rows, schema=raw_schema()), source)
    del rows
    gc.collect()
    typed = None
    filtered = None
    try:
        from pyspark import StorageLevel

        raw_frame = raw_frame.persist(StorageLevel.MEMORY_AND_DISK)
        missing, invalid = raw_quality(raw_frame)
        typed = typed_frame(raw_frame).persist(StorageLevel.MEMORY_AND_DISK)
        typed.write.mode("overwrite").parquet(str(SELECTED_DIR / source.periodo))
        raw_frame.unpersist()
        keys = key_audit(typed)
        quarter = value_counts(typed, "TRIMESTRE")
        year = value_counts(typed, "ANIO")
        codes = {column: value_counts(typed, column) for column in AUDITED_CODES}
        unrecognized = unrecognized_codes(typed)
        filtered, steps = apply_filters(typed)
        eligible = eligible_frame(filtered).persist(StorageLevel.MEMORY_AND_DISK)
        eligible.write.mode("overwrite").parquet(str(ELIGIBLE_PERIOD_DIR / source.periodo))
        eligible_count = eligible.count()
        unknown = unknown_counts(eligible)
        eligible.unpersist()
    finally:
        raw_frame.unpersist()
        if typed is not None:
            typed.unpersist()
        spark.catalog.clearCache()
        gc.collect()
    audit = {
        "periodo": source.periodo,
        "archivo": source.archivo,
        "sha256": digest,
        "bytes": source.path.stat().st_size,
        "mib": round(source.path.stat().st_size / MIB, 2),
        "hoja": extract.hoja,
        "hojas": extract.hojas,
        "hojas_visibles": extract.hojas_visibles,
        "columnas_encabezado": extract.columnas_encabezado,
        "encabezados_vacios": extract.encabezados_vacios,
        "encabezados_duplicados": extract.encabezados_duplicados,
        "filas_datos": extract.filas_datos,
        "filas_vacias_omitidas": extract.filas_vacias_omitidas,
        "columnas_requeridas": list(REQUIRED_COLUMNS),
        "anio_archivo": source.anio,
        "trimestre_calendario": source.trimestre,
        "faltantes_crudos": missing,
        "valores_no_convertibles": invalid,
        "trimestre_original": quarter,
        "anio_original_valores": year,
        "codigos_observados": codes,
        "codigos_no_reconocidos": unrecognized,
        "llaves": keys,
        "exclusiones": steps,
        "elegibles": eligible_count,
        "desconocidos_elegibles": unknown,
        "parquet_seleccion": relative(SELECTED_DIR / source.periodo),
        "parquet_elegibles": relative(ELIGIBLE_PERIOD_DIR / source.periodo),
        "segundos": round(time.perf_counter() - started, 1),
        "generado_utc": utc_now(),
    }
    write_json(PERIOD_AUDIT_DIR / f"{source.periodo}.json", audit)
    print(f"[{source.periodo}] elegibles={eligible_count} control={source.elegibles} ({audit['segundos']} s)")
    return audit


def union_periods(spark, base: Path, periods: list[str]):
    frames = [spark.read.parquet(str(base / period)) for period in periods]
    combined = frames[0]
    for frame in frames[1:]:
        combined = combined.unionByName(frame)
    return combined


def consolidate(spark) -> dict[str, Any]:
    from pyspark.sql import functions as F

    periods_2025 = [source.periodo for source in SOURCES if source.anio == 2025]
    periods_2026 = [source.periodo for source in SOURCES if source.anio == 2026]
    for periods, target in ((periods_2025, ELIGIBLE_2025), (periods_2026, ELIGIBLE_2026)):
        frame = union_periods(spark, ELIGIBLE_PERIOD_DIR, periods).orderBy("periodo_archivo", "fila_origen")
        frame.coalesce(1).write.mode("overwrite").parquet(str(target))
    selected = union_periods(spark, SELECTED_DIR, [source.periodo for source in SOURCES])
    total = selected.count()
    distinct_period_keys = selected.select("periodo_archivo", "NUM_HOGAR", "NUM_PERSONA").distinct().count()
    repeated_pairs = (
        selected.groupBy("NUM_HOGAR", "NUM_PERSONA")
        .agg(F.countDistinct("periodo_archivo").alias("periodos"))
        .filter(F.col("periodos") > 1)
        .count()
    )
    return {
        "filas_seleccionadas_totales": total,
        "llaves_periodo_distintas": distinct_period_keys,
        "llaves_periodo_duplicadas": total - distinct_period_keys,
        "pares_hogar_persona_en_varios_periodos": repeated_pairs,
    }


def dictionary_manifest() -> list[dict[str, Any]]:
    entries = []
    for source in SOURCES:
        path = source.dictionary_path
        present = path.exists()
        entries.append(
            {
                "periodo": source.periodo,
                "archivo": source.diccionario,
                "presente": present,
                "bytes": path.stat().st_size if present else None,
                "sha256": sha256_file(path) if present else None,
            }
        )
    return entries


def sample_records(spark, path: Path, limit: int = 5) -> list[dict[str, Any]]:
    rows = spark.read.parquet(str(path)).orderBy("periodo_archivo", "fila_origen").limit(limit).collect()
    return [row.asDict() for row in rows]


def run_foundation(force: bool, only: list[str] | None) -> int:
    started = time.perf_counter()
    missing_sources = [source.archivo for source in SOURCES if not source.path.exists()]
    if missing_sources:
        print(f"FALTAN FUENTES en {relative(DATA_DIR)}: {missing_sources}")
        return 2
    selected_sources = [source for source in SOURCES if only is None or source.periodo in only]
    FOUNDATION_OUT.mkdir(parents=True, exist_ok=True)
    spark = build_spark("lab7-foundation")
    spark.sparkContext.setLogLevel("WARN")
    checks = CheckLog()
    try:
        audits: dict[str, dict[str, Any]] = {}
        for source in SOURCES:
            digest = sha256_file(source.path)
            if source in selected_sources and (force or not period_is_current(source, digest)):
                audits[source.periodo] = process_source(spark, source, digest)
            elif period_is_current(source, digest):
                print(f"[{source.periodo}] reutilizando artefactos existentes (sha256 sin cambios)")
                audits[source.periodo] = read_json(PERIOD_AUDIT_DIR / f"{source.periodo}.json")
            else:
                print(f"[{source.periodo}] pendiente: ejecute --foundation sin --only o con --only {source.periodo}")
            gc.collect()
        if len(audits) != len(SOURCES):
            print("FUNDACION PARCIAL: faltan periodos por procesar; no se consolidan Parquet finales.")
            return 3
        for source in SOURCES:
            period_checks(audits[source.periodo], source, checks)
        cross = consolidate(spark)
        checks.require(cross["llaves_periodo_duplicadas"] == 0, "llave (periodo, hogar, persona) no unica")
        eligible_by_period = {period: audit["elegibles"] for period, audit in audits.items()}
        for name, (periods, expected) in CONTROL_TOTALS.items():
            observed = sum(eligible_by_period[period] for period in periods)
            checks.require(observed == expected, f"control {name}: {observed} != {expected}")
        raw_2025 = sum(audit["filas_datos"] for audit in audits.values() if audit["anio_archivo"] == 2025)
        raw_2026 = sum(audit["filas_datos"] for audit in audits.values() if audit["anio_archivo"] == 2026)
        checks.require(raw_2025 == RAW_TOTAL_CONTROLS["2025"], f"filas crudas 2025 {raw_2025}")
        checks.require(raw_2026 == RAW_TOTAL_CONTROLS["2026"], f"filas crudas 2026 {raw_2026}")
        manifest = {
            "generado_utc": utc_now(),
            "entorno": environment_info(),
            "spark": spark.version,
            "codigo_sha256": sha256_file(Path(__file__)),
            "columnas_requeridas": list(REQUIRED_COLUMNS),
            "fuentes": [
                {key: audits[source.periodo][key] for key in (
                    "periodo", "archivo", "sha256", "bytes", "mib", "hoja", "hojas", "hojas_visibles",
                    "columnas_encabezado", "filas_datos", "anio_archivo", "trimestre_calendario",
                )}
                for source in SOURCES
            ],
            "diccionarios": dictionary_manifest(),
            "limitaciones": [
                "Categorias conservadas como codigos normalizados; las etiquetas oficiales del diccionario se incorporaran en la fase de EDA.",
                "Rangos de codigos validados contra los codigos observados en la auditoria: DOMINIO 1-3, P03A03A 0-7, P05C16 1-9, OCUPADOS 1.",
                "periodo_archivo, anio_archivo y trimestre_calendario se derivan del manifiesto; TRIMESTRE original se conserva sin ajustes.",
            ],
            "salidas": {
                "eligible_2025": relative(ELIGIBLE_2025),
                "eligible_2026": relative(ELIGIBLE_2026),
                "seleccion_tipada": relative(SELECTED_DIR),
                "elegibles_por_periodo": relative(ELIGIBLE_PERIOD_DIR),
            },
        }
        write_json(FOUNDATION_OUT / "manifest.json", manifest)
        write_json(FOUNDATION_OUT / "audit_missing.json", {
            period: {"faltantes_crudos": audit["faltantes_crudos"], "valores_no_convertibles": audit["valores_no_convertibles"]}
            for period, audit in audits.items()
        })
        write_json(FOUNDATION_OUT / "audit_exclusions.json", {period: audit["exclusiones"] for period, audit in audits.items()})
        write_json(FOUNDATION_OUT / "audit_keys.json", {
            "por_periodo": {period: audit["llaves"] for period, audit in audits.items()},
            "entre_periodos": cross,
        })
        write_json(FOUNDATION_OUT / "audit_codes.json", {
            period: {
                "TRIMESTRE": audit["trimestre_original"],
                "ANIO": audit["anio_original_valores"],
                "codigos": audit["codigos_observados"],
                "no_reconocidos": audit["codigos_no_reconocidos"],
                "desconocidos_elegibles": audit["desconocidos_elegibles"],
            }
            for period, audit in audits.items()
        })
        write_json(FOUNDATION_OUT / "sample_eligible_2025.json", sample_records(spark, ELIGIBLE_2025))
        write_json(FOUNDATION_OUT / "schema_eligible.json", ELIGIBLE_SCHEMA)
        summary = {
            "generado_utc": utc_now(),
            "segundos": round(time.perf_counter() - started, 1),
            "elegibles_por_periodo": eligible_by_period,
            "controles": {name: expected for name, (_, expected) in CONTROL_TOTALS.items()},
            "filas_crudas": {"2025": raw_2025, "2026": raw_2026},
            "verificacion": checks.as_dict(),
        }
        write_json(FOUNDATION_OUT / "foundation_summary.json", summary)
        print_period_table(audits)
        status = "OK" if checks.ok else "FALLA"
        print(f"FUNDACION={status} controles={checks.passed} fallas={len(checks.failures)} advertencias={len(checks.warnings)}")
        return 0 if checks.ok else 1
    finally:
        spark.catalog.clearCache()
        spark.stop()


def print_period_table(audits: dict[str, dict[str, Any]]) -> None:
    print(f"{'periodo':<8}{'filas':>8}{'cols':>6}{'elegibles':>11}{'control':>9}")
    for source in SOURCES:
        audit = audits[source.periodo]
        print(
            f"{source.periodo:<8}{audit['filas_datos']:>8}{audit['columnas_encabezado']:>6}"
            f"{audit['elegibles']:>11}{source.elegibles:>9}"
        )


def run_preflight() -> int:
    checks = CheckLog()
    info = environment_info()
    print(f"Python {info['python']} | {info['plataforma']} | CPU {info['cpu']}")
    print(f"Java: {info['java']}")
    print(f"Paquetes: {info['paquetes']}")
    print(f"Memoria: {info['memoria']}")
    checks.require(sys.version_info[:2] >= (3, 10), "Python >= 3.10 requerido")
    checks.require(info["java"] is not None, "java no disponible en PATH")
    checks.require(info["ps_disponible"], "comando ps (procps) no disponible")
    for name in ("pyspark", "openpyxl", "pandas", "pyarrow", "numpy", "matplotlib"):
        checks.require(info["paquetes"][name] is not None, f"paquete {name} no instalado")
    pyspark_version = info["paquetes"]["pyspark"] or ""
    checks.require(pyspark_version.startswith("3.5."), f"PySpark 3.5.x requerido, encontrado {pyspark_version}")
    memory = info["memoria"]
    limit = memory["cgroup_limit_gib"] or memory["mem_total_gib"]
    checks.warn(
        limit is None or limit >= MIN_RECOMMENDED_MEMORY_GIB,
        f"memoria disponible para el contenedor {limit} GiB < {MIN_RECOMMENDED_MEMORY_GIB} GiB recomendados",
    )
    sources = []
    for source in SOURCES:
        present = source.path.exists()
        size = round(source.path.stat().st_size / MIB, 1) if present else None
        checks.require(present, f"falta {relative(source.path)}")
        if present:
            checks.warn(abs(size - source.tamano_mib) <= 0.2, f"{source.archivo}: {size} MiB vs {source.tamano_mib} MiB")
        dictionary_present = source.dictionary_path.exists()
        checks.require(dictionary_present, f"falta {relative(source.dictionary_path)}")
        sources.append({"periodo": source.periodo, "archivo": source.archivo, "presente": present, "mib": size,
                        "diccionario": source.diccionario, "diccionario_presente": dictionary_present})
        print(f"  {source.periodo}: {source.archivo} presente={present} MiB={size} diccionario={dictionary_present}")
    spark_result: dict[str, Any] = {}
    if info["paquetes"]["pyspark"] and info["java"]:
        spark = build_spark("lab7-preflight")
        spark.sparkContext.setLogLevel("WARN")
        try:
            frame = spark.range(10).selectExpr("id", "cast(id * 1.5 as double) as valor")
            frame.write.mode("overwrite").parquet(str(PREFLIGHT_PARQUET))
            restored = spark.read.parquet(str(PREFLIGHT_PARQUET))
            total = restored.agg({"id": "sum"}).first()[0]
            spark_result = {"version": spark.version, "filas": restored.count(), "suma_id": total}
            checks.require(spark.version.startswith("3.5."), f"Spark 3.5.x requerido, encontrado {spark.version}")
            checks.require(spark_result["filas"] == 10 and total == 45, "prueba Parquet de 10 filas fallida")
            print(f"Spark {spark.version}: Parquet 10 filas OK")
        finally:
            spark.stop()
            shutil.rmtree(PREFLIGHT_PARQUET, ignore_errors=True)
    write_json(FOUNDATION_OUT / "preflight.json", {
        "generado_utc": utc_now(),
        "entorno": info,
        "fuentes": sources,
        "spark": spark_result,
        "verificacion": checks.as_dict(),
    })
    status = "OK" if checks.ok else "FALLA"
    print(f"PREFLIGHT={status} controles={checks.passed} fallas={len(checks.failures)} advertencias={len(checks.warnings)}")
    return 0 if checks.ok else 1


def validate_eligible_frame(frame, label: str, expected_periods: list[str], checks: CheckLog) -> None:
    from pyspark.sql import functions as F

    schema = {item.name: item.dataType.simpleString() for item in frame.schema.fields}
    checks.require(schema == ELIGIBLE_SCHEMA, f"{label}: esquema distinto {schema}")
    counts = {row["periodo_archivo"]: int(row["count"]) for row in frame.groupBy("periodo_archivo").count().collect()}
    checks.require(sorted(counts) == sorted(expected_periods), f"{label}: periodos {sorted(counts)}")
    for period in expected_periods:
        expected = SOURCES_BY_PERIOD[period].elegibles
        checks.require(counts.get(period) == expected, f"{label} {period}: {counts.get(period)} != {expected}")
        print(f"  {label} {period}: {counts.get(period)} (control {expected})")
    total = frame.count()
    checks.require(frame.select("id_registro").distinct().count() == total, f"{label}: id_registro no unico")
    checks.require(
        frame.select("periodo_archivo", "NUM_HOGAR", "NUM_PERSONA").distinct().count() == total,
        f"{label}: llave (periodo, hogar, persona) no unica",
    )
    nulls = frame.agg(
        *[F.sum(F.when(F.col(column).isNull(), 1).otherwise(0)).alias(column) for column in MODEL_COLUMNS]
    ).first().asDict()
    checks.require(all((value or 0) == 0 for value in nulls.values()), f"{label}: nulos en variables de modelo {nulls}")
    violations = {
        "edad_menor_15": F.col("edad") < 15,
        "salario_no_positivo": F.col("salario_mensual") <= 0,
        "horas_fuera_rango": (F.col("horas_semanales") <= 0) | (F.col("horas_semanales") > 168),
        "antiguedad_negativa": F.col("antiguedad") < 0,
        "antiguedad_mayor_edad": F.col("antiguedad") > F.col("edad"),
        "ocupados_distinto_1": F.col("OCUPADOS") != 1,
        "categoria_fuera_1_4": ~F.col("categoria_ocupacional").isin("1", "2", "3", "4"),
        "nivel_educativo_fuera_0_7": ~F.col("nivel_educativo").isin(*[str(code) for code in range(8)]),
        "dominio_fuera_1_3": ~F.col("dominio").isin("1", "2", "3"),
        "anio_archivo_incoherente": F.col("anio_archivo") != F.lit(SOURCES_BY_PERIOD[expected_periods[0]].anio),
    }
    row = frame.agg(
        *[F.sum(F.when(condition, 1).otherwise(0)).alias(name) for name, condition in violations.items()]
    ).first().asDict()
    for name, count in row.items():
        checks.require((count or 0) == 0, f"{label}: {count} filas con {name}")
    for period in expected_periods:
        allowed = [int(key) for key in SOURCES_BY_PERIOD[period].trimestre_original]
        outside = frame.filter((F.col("periodo_archivo") == period) & ~F.col("TRIMESTRE").isin(*allowed)).count()
        checks.require(outside == 0, f"{label} {period}: TRIMESTRE original alterado")


def run_validate_foundation() -> int:
    checks = CheckLog()
    required_files = [FOUNDATION_OUT / name for name in (
        "manifest.json", "foundation_summary.json", "audit_missing.json", "audit_exclusions.json",
        "audit_keys.json", "audit_codes.json", "schema_eligible.json",
    )] + [PERIOD_AUDIT_DIR / f"{source.periodo}.json" for source in SOURCES]
    absent = [relative(path) for path in required_files if not path.exists()]
    absent += [relative(path) for path in (ELIGIBLE_2025, ELIGIBLE_2026) if not parquet_complete(path)]
    if absent:
        print(f"VALIDACION_FUNDACION=FALLA artefactos ausentes: {absent}")
        return 1
    print("Auditorias por periodo:")
    audits = {source.periodo: read_json(PERIOD_AUDIT_DIR / f"{source.periodo}.json") for source in SOURCES}
    for source in SOURCES:
        period_checks(audits[source.periodo], source, checks)
    manifest = read_json(FOUNDATION_OUT / "manifest.json")
    for entry in manifest["fuentes"]:
        checks.require(entry["sha256"] == audits[entry["periodo"]]["sha256"], f"{entry['periodo']}: hash de manifiesto")
    checks.require(all(item["presente"] for item in manifest["diccionarios"]), "diccionarios ausentes en manifiesto")
    keys = read_json(FOUNDATION_OUT / "audit_keys.json")["entre_periodos"]
    checks.require(keys["llaves_periodo_duplicadas"] == 0, "llaves (periodo, hogar, persona) duplicadas")
    checks.require(keys["filas_seleccionadas_totales"] == RAW_TOTAL_CONTROLS["total"], "total de filas crudas")
    print(f"  pares hogar-persona en varios periodos (longitudinal, se conservan): {keys['pares_hogar_persona_en_varios_periodos']}")
    spark = build_spark("lab7-validate-foundation")
    spark.sparkContext.setLogLevel("WARN")
    try:
        print("Parquet persistidos:")
        eligible_2025 = spark.read.parquet(str(ELIGIBLE_2025))
        eligible_2026 = spark.read.parquet(str(ELIGIBLE_2026))
        validate_eligible_frame(eligible_2025, "eligible_2025", [s.periodo for s in SOURCES if s.anio == 2025], checks)
        validate_eligible_frame(eligible_2026, "eligible_2026", [s.periodo for s in SOURCES if s.anio == 2026], checks)
        counts = {
            row["periodo_archivo"]: int(row["count"])
            for row in eligible_2025.unionByName(eligible_2026).groupBy("periodo_archivo").count().collect()
        }
        for name, (periods, expected) in CONTROL_TOTALS.items():
            observed = sum(counts.get(period, 0) for period in periods)
            checks.require(observed == expected, f"control {name}: {observed} != {expected}")
            print(f"  {name}: {observed} (control {expected})")
        overlap = eligible_2025.select("id_registro").intersect(eligible_2026.select("id_registro")).count()
        checks.require(overlap == 0, f"{overlap} id_registro compartidos entre 2025 y 2026")
    finally:
        spark.stop()
    result = {"generado_utc": utc_now(), **checks.as_dict()}
    write_json(FOUNDATION_OUT / "validation_foundation.json", result)
    status = "OK" if checks.ok else "FALLA"
    print(f"VALIDACION_FUNDACION={status} controles={checks.passed} fallas={len(checks.failures)} advertencias={len(checks.warnings)}")
    return 0 if checks.ok else 1


def category_label(column: str, code: Any) -> str:
    return CATEGORY_LABELS[column].get(str(code), UNKNOWN_LABEL)


def as_float(value: Any) -> float | None:
    if value is None:
        return None
    number = float(value)
    return number if math.isfinite(number) else None


def parquet_fingerprint(path: Path) -> dict[str, Any]:
    files = sorted(item for item in path.rglob("*.parquet") if item.is_file())
    digest = hashlib.sha256()
    for item in files:
        digest.update(item.name.encode("utf-8"))
        digest.update(sha256_file(item).encode("utf-8"))
    return {"archivos": len(files), "bytes": sum(item.stat().st_size for item in files), "sha256": digest.hexdigest()}


def load_eligible_2025(spark, checks: CheckLog):
    from pyspark.sql import functions as F

    frame = spark.read.parquet(str(ELIGIBLE_2025)).filter(F.col("anio_archivo") == 2025)
    counts = {row["periodo_archivo"]: int(row["count"]) for row in frame.groupBy("periodo_archivo").count().collect()}
    for source in SOURCES:
        if source.anio == 2025:
            checks.require(counts.get(source.periodo) == source.elegibles, f"{source.periodo}: {counts.get(source.periodo)} elegibles")
    expected_total = CONTROL_TOTALS["reentrenamiento_2025"][1]
    checks.require(sum(counts.values()) == expected_total, f"elegibles 2025 {sum(counts.values())} != {expected_total}")
    return frame


def percentile_expression(column: str):
    from pyspark.sql import functions as F

    levels = ", ".join(str(level) for level in PERCENTILES)
    return F.expr(f"percentile(`{column}`, array({levels}))")


def numeric_aggregations(column: str) -> list[Any]:
    from pyspark.sql import functions as F

    value = F.col(column)
    return [
        F.count(value).alias(f"{column}__n"),
        F.avg(value).alias(f"{column}__media"),
        F.stddev_samp(value).alias(f"{column}__desviacion"),
        F.min(value).alias(f"{column}__minimo"),
        F.max(value).alias(f"{column}__maximo"),
        percentile_expression(column).alias(f"{column}__percentiles"),
        F.skewness(value).alias(f"{column}__asimetria"),
        F.kurtosis(value).alias(f"{column}__curtosis"),
    ]


def summarize_numeric(frame, columns: tuple[str, ...], group_by: tuple[str, ...] = ()) -> list[dict[str, Any]]:
    aggregations = [expression for column in columns for expression in numeric_aggregations(column)]
    grouped = frame.groupBy(*group_by).agg(*aggregations).orderBy(*group_by) if group_by else frame.agg(*aggregations)
    summaries = []
    for row in grouped.collect():
        for column in columns:
            percentiles = row[f"{column}__percentiles"]
            item: dict[str, Any] = {key: row[key] for key in group_by}
            item.update(
                {
                    "variable": column,
                    "n": int(row[f"{column}__n"]),
                    "media": as_float(row[f"{column}__media"]),
                    "mediana": as_float(percentiles[1]),
                    "desviacion": as_float(row[f"{column}__desviacion"]),
                    "minimo": as_float(row[f"{column}__minimo"]),
                    "maximo": as_float(row[f"{column}__maximo"]),
                    "p25": as_float(percentiles[0]),
                    "p75": as_float(percentiles[2]),
                    "p95": as_float(percentiles[3]),
                    "asimetria": as_float(row[f"{column}__asimetria"]),
                    "curtosis": as_float(row[f"{column}__curtosis"]),
                }
            )
            summaries.append(item)
    return summaries


def category_table(frame, column: str, total: int) -> list[dict[str, Any]]:
    from pyspark.sql import functions as F

    rows = (
        frame.groupBy(column)
        .agg(
            F.count(F.lit(1)).alias("n"),
            F.avg("salario_mensual").alias("salario_media"),
            F.expr("percentile(salario_mensual, 0.5)").alias("salario_mediana"),
        )
        .collect()
    )
    table = [
        {
            "codigo": str(row[column]),
            "etiqueta": category_label(column, row[column]),
            "n": int(row["n"]),
            "porcentaje": 100.0 * int(row["n"]) / total,
            "salario_media": as_float(row["salario_media"]),
            "salario_mediana": as_float(row["salario_mediana"]),
        }
        for row in rows
    ]
    return sorted(table, key=lambda item: int(item["codigo"]) if item["codigo"].isdigit() else 99)


def composition_by_period(frame, column: str) -> list[dict[str, Any]]:
    from pyspark.sql import Window
    from pyspark.sql import functions as F

    window = Window.partitionBy("periodo_archivo")
    rows = (
        frame.groupBy("periodo_archivo", column)
        .agg(F.count(F.lit(1)).alias("n"), F.expr("percentile(salario_mensual, 0.5)").alias("salario_mediana"))
        .withColumn("porcentaje", 100.0 * F.col("n") / F.sum("n").over(window))
        .orderBy("periodo_archivo", column)
        .collect()
    )
    return [
        {
            "periodo_archivo": row["periodo_archivo"],
            "codigo": str(row[column]),
            "etiqueta": category_label(column, row[column]),
            "n": int(row["n"]),
            "porcentaje": float(row["porcentaje"]),
            "salario_mediana": as_float(row["salario_mediana"]),
        }
        for row in rows
    ]


def salary_asymmetry(frame, salary_summary: dict[str, Any]) -> dict[str, Any]:
    from pyspark.sql import functions as F

    p95 = salary_summary["p95"]
    row = frame.agg(
        F.skewness(F.log("salario_mensual")).alias("asimetria_log"),
        F.sum("salario_mensual").alias("masa_total"),
        F.sum(F.when(F.col("salario_mensual") >= p95, F.col("salario_mensual")).otherwise(0.0)).alias("masa_p95"),
        F.sum(F.when(F.col("salario_mensual") >= p95, 1).otherwise(0)).alias("n_p95"),
        F.sum(F.when(F.col("salario_mensual") > salary_summary["media"], 1).otherwise(0)).alias("n_sobre_media"),
    ).first()
    return {
        "asimetria": salary_summary["asimetria"],
        "curtosis_exceso": salary_summary["curtosis"],
        "razon_media_mediana": salary_summary["media"] / salary_summary["mediana"],
        "asimetria_log_salario": as_float(row["asimetria_log"]),
        "p95": p95,
        "n_mayor_igual_p95": int(row["n_p95"]),
        "porcentaje_masa_salarial_mayor_igual_p95": 100.0 * float(row["masa_p95"]) / float(row["masa_total"]),
        "porcentaje_sobre_media": 100.0 * int(row["n_sobre_media"]) / salary_summary["n"],
    }


def pearson_matrix(frame) -> list[list[float]]:
    from pyspark.ml.feature import VectorAssembler
    from pyspark.ml.stat import Correlation

    assembler = VectorAssembler(inputCols=list(NUMERIC_VARIABLES), outputCol="vector_correlacion")
    vectors = assembler.transform(frame.select(*NUMERIC_VARIABLES)).select("vector_correlacion")
    matrix = Correlation.corr(vectors, "vector_correlacion", "pearson").first()[0].toArray()
    return [[float(value) for value in row] for row in matrix]


def cluster_input(frame):
    from pyspark.sql import functions as F

    return frame.withColumn("nivel_educativo_ordinal", F.col("nivel_educativo").cast("double"))


def fit_standardizer(frame):
    from pyspark.ml import Pipeline
    from pyspark.ml.feature import StandardScaler, VectorAssembler

    assembler = VectorAssembler(inputCols=list(CLUSTER_FEATURES), outputCol="features_raw")
    scaler = StandardScaler(inputCol="features_raw", outputCol="features_std", withMean=True, withStd=True)
    return Pipeline(stages=[assembler, scaler]).fit(frame)


def compare_kmeans(scaled) -> tuple[list[dict[str, Any]], dict[int, Any]]:
    from pyspark.ml.clustering import KMeans
    from pyspark.ml.evaluation import ClusteringEvaluator

    evaluator = ClusteringEvaluator(
        featuresCol="features_std", predictionCol="cluster", metricName="silhouette", distanceMeasure="squaredEuclidean"
    )
    results = []
    models = {}
    for k in K_RANGE:
        started = time.perf_counter()
        model = KMeans(
            featuresCol="features_std", predictionCol="cluster", k=k, seed=SEED, maxIter=KMEANS_MAX_ITER, initMode="k-means||"
        ).fit(scaled)
        silhouette = evaluator.evaluate(model.transform(scaled))
        results.append(
            {
                "k": k,
                "silueta": float(silhouette),
                "wssse": float(model.summary.trainingCost),
                "tamanos": [int(size) for size in model.summary.clusterSizes],
                "iteraciones": int(model.summary.numIter),
                "segundos": round(time.perf_counter() - started, 1),
            }
        )
        models[k] = model
        print(f"  K={k}: silueta={silhouette:.4f} WSSSE={model.summary.trainingCost:,.1f}")
    for previous, current in zip(results, results[1:]):
        current["reduccion_wssse_pct"] = 100.0 * (previous["wssse"] - current["wssse"]) / previous["wssse"]
    return results, models


def select_k(results: list[dict[str, Any]]) -> int:
    return max(results, key=lambda item: (round(item["silueta"], 6), -item["k"]))["k"]


def top_category(composition: list[dict[str, Any]]) -> dict[str, Any]:
    return max(composition, key=lambda item: item["n"])


def cluster_description(profile: dict[str, Any], global_medians: dict[str, float]) -> str:
    stats = profile["estadisticas"]
    education = top_category(profile["composicion"]["nivel_educativo"])
    category = top_category(profile["composicion"]["categoria_ocupacional"])
    domain = top_category(profile["composicion"]["dominio"])
    salary_ratio = stats["salario_mensual"]["mediana"] / global_medians["salario_mensual"]
    return (
        f"Cluster {profile['cluster']}: {profile['n']:,} registros ({profile['porcentaje']:.1f}%). "
        f"Edad mediana {stats['edad']['mediana']:.0f} años (global {global_medians['edad']:.0f}), "
        f"antigüedad mediana {stats['antiguedad']['mediana']:.1f} años (global {global_medians['antiguedad']:.1f}) "
        f"y {stats['horas_semanales']['mediana']:.0f} horas semanales medianas (global {global_medians['horas_semanales']:.0f}). "
        f"Nivel educativo más frecuente: {education['etiqueta']} ({education['porcentaje']:.1f}%); "
        f"categoría más frecuente: {category['etiqueta']} ({category['porcentaje']:.1f}%); "
        f"dominio más frecuente: {domain['etiqueta']} ({domain['porcentaje']:.1f}%). "
        f"Salario mensual mediano Q{stats['salario_mensual']['mediana']:,.0f} y medio Q{stats['salario_mensual']['media']:,.0f}, "
        f"{salary_ratio:.2f} veces la mediana global; el salario no participó en el ajuste del cluster."
    )


def cluster_profiles(assigned, model, standardizer, global_medians: dict[str, float]) -> list[dict[str, Any]]:
    from pyspark.sql import functions as F

    total = assigned.count()
    numeric = summarize_numeric(assigned, NUMERIC_VARIABLES, ("cluster",))
    scaler = standardizer.stages[-1]
    means = scaler.mean.toArray()
    deviations = scaler.std.toArray()
    centers = model.clusterCenters()
    profiles = []
    for cluster in range(len(centers)):
        size = assigned.filter(F.col("cluster") == cluster).count()
        stats = {
            item["variable"]: {key: item[key] for key in ("n", "media", "mediana", "desviacion", "p25", "p75", "minimo", "maximo")}
            for item in numeric
            if item["cluster"] == cluster
        }
        composition = {}
        for column in CATEGORY_LABELS:
            subset = assigned.filter(F.col("cluster") == cluster)
            table = category_table(subset, column, max(size, 1))
            composition[column] = sorted(table, key=lambda item: -item["n"])
        center = centers[cluster]
        profile = {
            "cluster": cluster,
            "n": size,
            "porcentaje": 100.0 * size / total,
            "estadisticas": stats,
            "composicion": composition,
            "centroide_estandarizado": {feature: float(center[index]) for index, feature in enumerate(CLUSTER_FEATURES)},
            "centroide_unidades_originales": {
                feature: float(center[index] * deviations[index] + means[index]) for index, feature in enumerate(CLUSTER_FEATURES)
            },
        }
        profile["descripcion"] = cluster_description(profile, global_medians)
        profiles.append(profile)
    return sorted(profiles, key=lambda item: item["estadisticas"]["salario_mensual"]["mediana"])


def deterministic_sample(frame, size: int = SAMPLE_SIZE):
    from pyspark.sql import functions as F

    ordered = frame.withColumn("_orden_muestra", F.xxhash64(F.col("id_registro"), F.lit(SEED)))
    return ordered.orderBy("_orden_muestra", "id_registro").limit(size).drop("_orden_muestra")


def collect_pandas(frame, columns: list[str]):
    import pandas as pd

    rows = frame.select(*columns).collect()
    return pd.DataFrame([row.asDict() for row in rows], columns=columns)


def compute_eda(spark, frame) -> dict[str, Any]:
    started = time.perf_counter()
    total = frame.count()
    print(f"EDA sobre {total:,} registros elegibles de 2025")
    descriptives = summarize_numeric(frame, NUMERIC_VARIABLES)
    by_name = {item["variable"]: item for item in descriptives}
    by_period = summarize_numeric(frame, NUMERIC_VARIABLES, ("periodo_archivo",))
    categories = {column: category_table(frame, column, total) for column in CATEGORY_LABELS}
    composition = {column: composition_by_period(frame, column) for column in CATEGORY_LABELS}
    asymmetry = salary_asymmetry(frame, by_name["salario_mensual"])
    print("Correlacion de Pearson con Spark")
    correlation = pearson_matrix(frame)
    print("KMeans estandarizado sin salario")
    clustering_frame = cluster_input(frame)
    standardizer = fit_standardizer(clustering_frame)
    scaled = standardizer.transform(clustering_frame).cache()
    try:
        comparison, models = compare_kmeans(scaled)
        chosen = select_k(comparison)
        assigned = models[chosen].transform(scaled).drop("features_raw", "features_std").cache()
        global_medians = {name: by_name[name]["mediana"] for name in NUMERIC_VARIABLES}
        profiles = cluster_profiles(assigned, models[chosen], standardizer, global_medians)
    finally:
        scaled.unpersist()
    return {
        "n": total,
        "descriptivos": descriptives,
        "descriptivos_trimestrales": by_period,
        "categorias": categories,
        "composicion_trimestral": composition,
        "asimetria_salario": asymmetry,
        "correlacion": {"variables": list(NUMERIC_VARIABLES), "metodo": "pearson", "n": total, "matriz": correlation},
        "kmeans": {
            "resultados": comparison,
            "k_seleccionado": chosen,
            "en_limite_superior": chosen == max(K_RANGE),
            "criterio": "mayor silueta (distancia euclidiana cuadrada) entre K=2..5; WSSSE se reporta como evidencia complementaria de codo",
            "alcance": "mejor K entre los valores probados, no un optimo universal",
            "variables_ajuste": list(CLUSTER_FEATURES),
            "variables_caracterizacion": ["salario_mensual", "categoria_ocupacional", "dominio"],
            "salario_excluido_del_ajuste": True,
            "estandarizacion": "StandardScaler(withMean=True, withStd=True) ajustado con todos los elegibles de 2025",
            "nivel_educativo": "codigo ordinal 0-7 tratado como numerico para el ajuste",
            "semilla": SEED,
            "max_iter": KMEANS_MAX_ITER,
            "init_mode": "k-means||",
        },
        "perfiles": profiles,
        "medianas_globales": global_medians,
        "_modelos": models,
        "_estandarizador": standardizer,
        "_asignados": assigned,
        "segundos_calculo": round(time.perf_counter() - started, 1),
    }


def save_figure(figure, path: Path) -> str:
    path.parent.mkdir(parents=True, exist_ok=True)
    figure.tight_layout()
    figure.savefig(path, dpi=130)
    import matplotlib.pyplot as plt

    plt.close(figure)
    return relative(path)


def figure_distributions(sample, path: Path) -> str:
    import matplotlib.pyplot as plt
    import numpy as np

    figure, axes = plt.subplots(2, 2, figsize=(11, 7.5))
    for axis, column in zip(axes.ravel(), NUMERIC_VARIABLES):
        values = sample[column].astype(float)
        if column == "salario_mensual":
            bins = np.logspace(np.log10(values.min()), np.log10(values.max()), 50)
            axis.hist(values, bins=bins, color="#4C72B0", edgecolor="white")
            axis.set_xscale("log")
        else:
            axis.hist(values, bins=40, color="#4C72B0", edgecolor="white")
        axis.axvline(values.median(), color="#C44E52", linestyle="--", label="mediana muestra")
        axis.set_title(NUMERIC_TITLES[column])
        axis.set_ylabel("Registros")
        axis.legend(fontsize=8)
    figure.suptitle(f"Distribuciones 2025, muestra determinista de {len(sample):,} registros")
    return save_figure(figure, path)


def figure_salary_by_category(categories: dict[str, list[dict[str, Any]]], path: Path) -> str:
    import matplotlib.pyplot as plt

    figure, axes = plt.subplots(1, 3, figsize=(15, 5))
    for axis, column in zip(axes, CATEGORY_LABELS):
        table = categories[column]
        labels = [f"{item['etiqueta']} (n={item['n']:,})" for item in table]
        axis.barh(labels, [item["salario_mediana"] for item in table], color="#55A868")
        axis.set_title(f"Salario mediano por {CATEGORY_TITLES[column].lower()}")
        axis.set_xlabel("Q mensuales")
        axis.invert_yaxis()
    return save_figure(figure, path)


def figure_quarterly(by_period: list[dict[str, Any]], composition: list[dict[str, Any]], path: Path) -> str:
    import matplotlib.pyplot as plt
    import pandas as pd

    salary = [item for item in by_period if item["variable"] == "salario_mensual"]
    periods = [item["periodo_archivo"] for item in salary]
    figure, axes = plt.subplots(1, 2, figsize=(13, 4.8))
    axes[0].plot(periods, [item["media"] for item in salary], marker="o", label="media")
    axes[0].plot(periods, [item["mediana"] for item in salary], marker="s", label="mediana")
    axes[0].fill_between(periods, [item["p25"] for item in salary], [item["p75"] for item in salary], alpha=0.15, label="P25-P75")
    axes[0].set_title("Salario mensual por trimestre (2025)")
    axes[0].set_ylabel("Q mensuales")
    axes[0].legend()
    shares = pd.DataFrame(composition).pivot(index="periodo_archivo", columns="etiqueta", values="porcentaje").fillna(0.0)
    bottom = [0.0] * len(shares)
    for label in shares.columns:
        axes[1].bar(shares.index, shares[label], bottom=bottom, label=label)
        bottom = [base + value for base, value in zip(bottom, shares[label])]
    axes[1].set_title("Composición por categoría ocupacional")
    axes[1].set_ylabel("% de registros elegibles")
    axes[1].legend(fontsize=8, loc="lower right")
    return save_figure(figure, path)


def figure_heatmap(matrix: list[list[float]], rows: list[str], columns: list[str], title: str, path: Path, symmetric: bool) -> str:
    import matplotlib.pyplot as plt
    import numpy as np

    values = np.array(matrix)
    limit = 1.0 if symmetric else max(abs(values.min()), abs(values.max()))
    figure, axis = plt.subplots(figsize=(1.6 * len(columns) + 3, 1.0 * len(rows) + 2))
    image = axis.imshow(values, cmap="RdBu_r", vmin=-limit, vmax=limit)
    axis.set_xticks(range(len(columns)), columns, rotation=30, ha="right")
    axis.set_yticks(range(len(rows)), rows)
    for i in range(values.shape[0]):
        for j in range(values.shape[1]):
            axis.text(j, i, f"{values[i, j]:.2f}", ha="center", va="center", fontsize=9)
    figure.colorbar(image, ax=axis, shrink=0.8)
    axis.set_title(title)
    return save_figure(figure, path)


def figure_kmeans_selection(comparison: list[dict[str, Any]], chosen: int, path: Path) -> str:
    import matplotlib.pyplot as plt

    ks = [item["k"] for item in comparison]
    figure, axes = plt.subplots(1, 2, figsize=(11, 4.2))
    axes[0].plot(ks, [item["silueta"] for item in comparison], marker="o")
    axes[0].set_title("Silueta por K")
    axes[1].plot(ks, [item["wssse"] for item in comparison], marker="o", color="#C44E52")
    axes[1].set_title("WSSSE por K")
    for axis in axes:
        axis.axvline(chosen, color="gray", linestyle="--", label=f"K elegido = {chosen}")
        axis.set_xticks(ks)
        axis.set_xlabel("K")
        axis.legend()
    return save_figure(figure, path)


def figure_cluster_scatter(sample, path: Path) -> str:
    import matplotlib.pyplot as plt

    figure, axes = plt.subplots(1, 2, figsize=(13, 5))
    for cluster in sorted(sample["cluster"].unique()):
        subset = sample[sample["cluster"] == cluster]
        axes[0].scatter(subset["edad"], subset["salario_mensual"], s=6, alpha=0.45, label=f"cluster {cluster}")
        axes[1].scatter(subset["antiguedad"], subset["horas_semanales"], s=6, alpha=0.45, label=f"cluster {cluster}")
    axes[0].set_yscale("log")
    axes[0].set_xlabel("Edad (años)")
    axes[0].set_ylabel("Salario mensual (Q, escala log)")
    axes[1].set_xlabel("Antigüedad (años)")
    axes[1].set_ylabel("Horas semanales")
    axes[0].legend(markerscale=3, fontsize=8)
    figure.suptitle(f"Clusters en muestra determinista de {len(sample):,} registros")
    return save_figure(figure, path)


def write_eda_figures(results: dict[str, Any], sample) -> dict[str, str]:
    profiles = sorted(results["perfiles"], key=lambda item: item["cluster"])
    return {
        "distribuciones": figure_distributions(sample, EDA_FIGURES / "distribuciones_2025.png"),
        "salario_por_categoria": figure_salary_by_category(results["categorias"], EDA_FIGURES / "salario_mediano_por_categoria.png"),
        "evolucion_trimestral": figure_quarterly(
            results["descriptivos_trimestrales"],
            results["composicion_trimestral"]["categoria_ocupacional"],
            EDA_FIGURES / "evolucion_trimestral_2025.png",
        ),
        "correlacion": figure_heatmap(
            results["correlacion"]["matriz"],
            [NUMERIC_TITLES[name] for name in NUMERIC_VARIABLES],
            [NUMERIC_TITLES[name] for name in NUMERIC_VARIABLES],
            "Correlación de Pearson, elegibles 2025",
            EDA_FIGURES / "correlacion_pearson_2025.png",
            True,
        ),
        "kmeans_seleccion": figure_kmeans_selection(
            results["kmeans"]["resultados"], results["kmeans"]["k_seleccionado"], EDA_FIGURES / "kmeans_silueta_wssse.png"
        ),
        "clusters_centroides": figure_heatmap(
            [[profile["centroide_estandarizado"][feature] for profile in profiles] for feature in CLUSTER_FEATURES],
            list(CLUSTER_FEATURES),
            [f"cluster {profile['cluster']}" for profile in profiles],
            "Centroides estandarizados (z)",
            EDA_FIGURES / "kmeans_centroides_estandarizados.png",
            False,
        ),
        "clusters_muestra": figure_cluster_scatter(sample, EDA_FIGURES / "kmeans_clusters_muestra.png"),
    }


def label_coverage(results: dict[str, Any]) -> list[str]:
    return [
        f"{column}={item['codigo']}"
        for column, table in results["categorias"].items()
        for item in table
        if item["etiqueta"] == UNKNOWN_LABEL
    ]


def eda_tables(results: dict[str, Any]) -> dict[str, Any]:
    return {
        "descriptivos_2025.json": {
            "n": results["n"],
            "descriptivos": results["descriptivos"],
            "asimetria_salario": results["asimetria_salario"],
        },
        "descriptivos_trimestrales_2025.json": {
            "numericos": results["descriptivos_trimestrales"],
            "composicion": results["composicion_trimestral"],
        },
        "categorias_2025.json": {"etiquetas": CATEGORY_LABELS, "tablas": results["categorias"]},
        "correlacion_pearson_2025.json": results["correlacion"],
        "kmeans_seleccion.json": results["kmeans"],
        "clusters_perfiles.json": {
            "k": results["kmeans"]["k_seleccionado"],
            "medianas_globales": results["medianas_globales"],
            "perfiles": results["perfiles"],
        },
    }


def run_eda_clustering() -> int:
    started = time.perf_counter()
    checks = CheckLog()
    validation_path = FOUNDATION_OUT / "validation_foundation.json"
    foundation_ok = validation_path.exists() and read_json(validation_path).get("ok") is True and parquet_complete(ELIGIBLE_2025)
    if not foundation_ok:
        print("EDA_CLUSTERING=FALLA la fundacion no esta validada: ejecute --validate-foundation")
        return 1
    EDA_OUT.mkdir(parents=True, exist_ok=True)
    spark = build_spark("lab7-eda-clustering")
    spark.sparkContext.setLogLevel("WARN")
    frame = None
    assigned = None
    try:
        frame = load_eligible_2025(spark, checks).cache()
        if not checks.ok:
            print("EDA_CLUSTERING=FALLA conteos de entrada distintos de los controles")
            return 1
        results = compute_eda(spark, frame)
        assigned = results.pop("_asignados")
        models = results.pop("_modelos")
        standardizer = results.pop("_estandarizador")
        unknown = label_coverage(results)
        checks.require(not unknown, f"codigos sin etiqueta oficial: {unknown}")
        from pyspark.ml import PipelineModel

        chosen = results["kmeans"]["k_seleccionado"]
        pipeline = PipelineModel(stages=[*standardizer.stages, models[chosen]])
        pipeline.write().overwrite().save(str(KMEANS_MODEL))
        assigned.select("id_registro", "periodo_archivo", "cluster").orderBy("periodo_archivo", "fila_origen").coalesce(1).write.mode(
            "overwrite"
        ).parquet(str(CLUSTERS_2025))
        sample = collect_pandas(
            deterministic_sample(assigned), ["id_registro", *NUMERIC_VARIABLES, *CATEGORY_LABELS, "cluster"]
        )
        figures = write_eda_figures(results, sample)
        for name, payload in eda_tables(results).items():
            write_json(EDA_OUT / name, payload)
        summary = {
            "generado_utc": utc_now(),
            "segundos": round(time.perf_counter() - started, 1),
            "spark": spark.version,
            "codigo_sha256": sha256_file(Path(__file__)),
            "entrada": {"ruta": relative(ELIGIBLE_2025), "n": results["n"], **parquet_fingerprint(ELIGIBLE_2025)},
            "muestra_graficos": {
                "n": len(sample),
                "metodo": f"orden por xxhash64(id_registro, {SEED}) y limite {SAMPLE_SIZE}",
                "uso": "solo histogramas y dispersiones; estadisticas y metricas con Spark completo",
            },
            "k_seleccionado": chosen,
            "salidas": {
                "tablas": [relative(EDA_OUT / name) for name in eda_tables(results)],
                "figuras": figures,
                "clusters": relative(CLUSTERS_2025),
                "modelo_kmeans": relative(KMEANS_MODEL),
            },
            "verificacion": checks.as_dict(),
        }
        write_json(EDA_OUT / "eda_summary.json", summary)
        salary = next(item for item in results["descriptivos"] if item["variable"] == "salario_mensual")
        print(f"Salario 2025: media Q{salary['media']:,.2f} mediana Q{salary['mediana']:,.2f} P95 Q{salary['p95']:,.2f}")
        for item in results["kmeans"]["resultados"]:
            print(f"  K={item['k']} silueta={item['silueta']:.4f} WSSSE={item['wssse']:,.1f} tamanos={item['tamanos']}")
        print(f"K seleccionado={chosen} (mejor entre K=2..5; limite superior={results['kmeans']['en_limite_superior']})")
        status = "OK" if checks.ok else "FALLA"
        print(f"EDA_CLUSTERING={status} controles={checks.passed} fallas={len(checks.failures)} figuras={len(figures)}")
        return 0 if checks.ok else 1
    finally:
        if assigned is not None:
            assigned.unpersist()
        if frame is not None:
            frame.unpersist()
        spark.catalog.clearCache()
        spark.stop()


def close_to(left: float, right: float, tolerance: float = 1e-6) -> bool:
    return abs(left - right) <= tolerance * max(1.0, abs(left), abs(right))


def validate_eda_tables(tables: dict[str, Any], expected_total: int, checks: CheckLog) -> None:
    descriptives = tables["descriptivos_2025.json"]
    checks.require(descriptives["n"] == expected_total, f"n EDA {descriptives['n']} != {expected_total}")
    checks.require({item["variable"] for item in descriptives["descriptivos"]} == set(NUMERIC_VARIABLES), "variables descriptivas")
    for item in descriptives["descriptivos"]:
        name = item["variable"]
        checks.require(item["n"] == expected_total, f"{name}: n {item['n']}")
        ordered = [item["minimo"], item["p25"], item["mediana"], item["p75"], item["p95"], item["maximo"]]
        checks.require(all(a <= b for a, b in zip(ordered, ordered[1:])), f"{name}: percentiles no ordenados")
        checks.require(item["minimo"] <= item["media"] <= item["maximo"], f"{name}: media fuera de rango")
        checks.require(item["desviacion"] is not None and item["desviacion"] >= 0, f"{name}: desviacion invalida")
        checks.require(item["asimetria"] is not None, f"{name}: asimetria ausente")
    salary = next(item for item in descriptives["descriptivos"] if item["variable"] == "salario_mensual")
    checks.require(salary["minimo"] > 0, "salario minimo no positivo")
    asymmetry = descriptives["asimetria_salario"]
    checks.require(asymmetry["n_mayor_igual_p95"] >= 1, "banda salarial >= P95 vacia")
    checks.require(close_to(asymmetry["p95"], salary["p95"]), "P95 inconsistente")
    periods_2025 = sorted(source.periodo for source in SOURCES if source.anio == 2025)
    quarterly = tables["descriptivos_trimestrales_2025.json"]
    salary_by_period = {item["periodo_archivo"]: item["n"] for item in quarterly["numericos"] if item["variable"] == "salario_mensual"}
    checks.require(sorted(salary_by_period) == periods_2025, f"periodos trimestrales {sorted(salary_by_period)}")
    for period, n in salary_by_period.items():
        checks.require(n == SOURCES_BY_PERIOD[period].elegibles, f"{period}: n trimestral {n}")
    for column, rows in quarterly["composicion"].items():
        for period in periods_2025:
            share = sum(item["porcentaje"] for item in rows if item["periodo_archivo"] == period)
            checks.require(close_to(share, 100.0, 1e-6), f"{column} {period}: composicion suma {share}")
    categories = tables["categorias_2025.json"]["tablas"]
    for column, table in categories.items():
        checks.require(sum(item["n"] for item in table) == expected_total, f"{column}: conteos no suman el total")
        checks.require(all(item["etiqueta"] == CATEGORY_LABELS[column].get(item["codigo"]) for item in table), f"{column}: etiquetas")
        checks.require(all(item["salario_mediana"] > 0 for item in table), f"{column}: salario mediano invalido")
    correlation = tables["correlacion_pearson_2025.json"]
    matrix = correlation["matriz"]
    size = len(NUMERIC_VARIABLES)
    checks.require(correlation["variables"] == list(NUMERIC_VARIABLES) and correlation["metodo"] == "pearson", "correlacion: variables")
    checks.require(correlation["n"] == expected_total, "correlacion: n")
    checks.require(len(matrix) == size and all(len(row) == size for row in matrix), "correlacion: dimension")
    for i in range(size):
        checks.require(close_to(matrix[i][i], 1.0), f"correlacion: diagonal {i}")
        for j in range(size):
            checks.require(-1.0 - 1e-9 <= matrix[i][j] <= 1.0 + 1e-9, f"correlacion: rango {i},{j}")
            checks.require(close_to(matrix[i][j], matrix[j][i]), f"correlacion: simetria {i},{j}")
    kmeans = tables["kmeans_seleccion.json"]
    results = kmeans["resultados"]
    checks.require([item["k"] for item in results] == list(K_RANGE), "kmeans: K evaluados")
    checks.require("salario_mensual" not in kmeans["variables_ajuste"] and kmeans["salario_excluido_del_ajuste"], "kmeans: salario en ajuste")
    checks.require(kmeans["variables_ajuste"] == list(CLUSTER_FEATURES), "kmeans: variables de ajuste")
    checks.require(kmeans["semilla"] == SEED, "kmeans: semilla")
    for item in results:
        checks.require(-1.0 <= item["silueta"] <= 1.0, f"K={item['k']}: silueta fuera de rango")
        checks.require(item["wssse"] > 0, f"K={item['k']}: WSSSE no positivo")
        checks.require(len(item["tamanos"]) == item["k"] and sum(item["tamanos"]) == expected_total, f"K={item['k']}: tamanos")
        checks.require(min(item["tamanos"]) > 0, f"K={item['k']}: cluster vacio")
    for previous, current in zip(results, results[1:]):
        checks.warn(current["wssse"] <= previous["wssse"], f"WSSSE aumenta de K={previous['k']} a K={current['k']}")
    chosen = kmeans["k_seleccionado"]
    checks.require(chosen == select_k(results), "kmeans: K elegido no coincide con la mayor silueta")
    checks.require(kmeans["en_limite_superior"] == (chosen == max(K_RANGE)), "kmeans: bandera de limite")
    checks.require("no un optimo universal" in kmeans["alcance"], "kmeans: alcance del K no declarado")
    profiles = tables["clusters_perfiles.json"]
    checks.require(profiles["k"] == chosen and len(profiles["perfiles"]) == chosen, "perfiles: numero de clusters")
    chosen_sizes = sorted(next(item["tamanos"] for item in results if item["k"] == chosen))
    checks.require(sorted(profile["n"] for profile in profiles["perfiles"]) == chosen_sizes, "perfiles: tamanos distintos de KMeans")
    for profile in profiles["perfiles"]:
        label = f"cluster {profile['cluster']}"
        checks.require(set(profile["estadisticas"]) == set(NUMERIC_VARIABLES), f"{label}: estadisticas")
        checks.require(all(stats["n"] == profile["n"] for stats in profile["estadisticas"].values()), f"{label}: n por variable")
        for column in CATEGORY_LABELS:
            checks.require(sum(item["n"] for item in profile["composicion"][column]) == profile["n"], f"{label}: composicion {column}")
        checks.require(set(profile["centroide_estandarizado"]) == set(CLUSTER_FEATURES), f"{label}: centroide")
        checks.require(
            profile["descripcion"].startswith(f"Cluster {profile['cluster']}:") and "Q" in profile["descripcion"], f"{label}: descripcion"
        )


def validate_figures(figures: dict[str, str], checks: CheckLog) -> None:
    checks.require(len(figures) >= 7, f"figuras registradas {len(figures)}")
    for name, path in figures.items():
        file = ROOT / path
        valid = file.exists() and file.stat().st_size > 5_000 and file.read_bytes()[:8] == b"\x89PNG\r\n\x1a\n"
        checks.require(valid, f"figura {name} ausente o invalida: {path}")


def run_validate_eda_clustering() -> int:
    checks = CheckLog()
    names = [*eda_tables_names(), "eda_summary.json"]
    absent = [relative(EDA_OUT / name) for name in names if not (EDA_OUT / name).exists()]
    absent += [relative(path) for path in (CLUSTERS_2025,) if not parquet_complete(path)]
    absent += [relative(KMEANS_MODEL)] if not (KMEANS_MODEL / "metadata").exists() else []
    if absent:
        print(f"VALIDACION_EDA_CLUSTERING=FALLA artefactos ausentes: {absent}")
        return 1
    expected_total = CONTROL_TOTALS["reentrenamiento_2025"][1]
    tables = {name: read_json(EDA_OUT / name) for name in eda_tables_names()}
    summary = read_json(EDA_OUT / "eda_summary.json")
    print("Tablas y metricas persistidas:")
    validate_eda_tables(tables, expected_total, checks)
    checks.require(summary["verificacion"]["ok"], "eda_summary registra fallas")
    checks.require(summary["entrada"]["n"] == expected_total, "eda_summary: n de entrada")
    checks.require(0 < summary["muestra_graficos"]["n"] <= SAMPLE_SIZE, "muestra de graficos fuera de limite")
    checks.require(summary["entrada"]["sha256"] == parquet_fingerprint(ELIGIBLE_2025)["sha256"], "eligible_2025 cambio despues del EDA")
    validate_figures(summary["salidas"]["figuras"], checks)
    chosen = tables["kmeans_seleccion.json"]["k_seleccionado"]
    spark = build_spark("lab7-validate-eda-clustering")
    spark.sparkContext.setLogLevel("WARN")
    try:
        from pyspark.sql import functions as F

        clusters = spark.read.parquet(str(CLUSTERS_2025))
        eligible = spark.read.parquet(str(ELIGIBLE_2025))
        total = clusters.count()
        checks.require(total == expected_total, f"clusters_2025: {total} filas")
        checks.require(clusters.select("id_registro").distinct().count() == total, "clusters_2025: id_registro duplicado")
        orphan = clusters.join(eligible.select("id_registro"), "id_registro", "left_anti").count()
        checks.require(orphan == 0, f"clusters_2025: {orphan} ids fuera de eligible_2025")
        sizes = sorted(int(row["count"]) for row in clusters.groupBy("cluster").count().collect())
        expected_sizes = sorted(profile["n"] for profile in tables["clusters_perfiles.json"]["perfiles"])
        checks.require(sizes == expected_sizes, f"clusters_2025: tamanos {sizes} != {expected_sizes}")
        checks.require(clusters.filter(F.col("cluster").isNull()).count() == 0, "clusters_2025: cluster nulo")
        print(f"  clusters_2025: {total:,} filas, K={chosen}, tamanos={sizes}")
    finally:
        spark.stop()
    write_json(EDA_OUT / "validation_eda_clustering.json", {"generado_utc": utc_now(), **checks.as_dict()})
    status = "OK" if checks.ok else "FALLA"
    print(f"VALIDACION_EDA_CLUSTERING={status} controles={checks.passed} fallas={len(checks.failures)} advertencias={len(checks.warnings)}")
    return 0 if checks.ok else 1


def eda_tables_names() -> list[str]:
    return [
        "descriptivos_2025.json",
        "descriptivos_trimestrales_2025.json",
        "categorias_2025.json",
        "correlacion_pearson_2025.json",
        "kmeans_seleccion.json",
        "clusters_perfiles.json",
    ]


def temporal_partitions(spark, checks: CheckLog):
    from pyspark.sql import functions as F

    frame = spark.read.parquet(str(ELIGIBLE_2025))
    foreign = frame.filter((F.col("anio_archivo") != 2025) | F.col("periodo_archivo").startswith("2026")).count()
    checks.require(foreign == 0, f"{foreign} filas ajenas a 2025 en {relative(ELIGIBLE_2025)}")
    columns = ["id_registro", "periodo_archivo", "fila_origen", TARGET, *NUMERIC_PREDICTORS, *CATEGORICAL_PREDICTORS]
    selected = frame.select(*columns)
    train = selected.filter(F.col("periodo_archivo").isin(*TRAIN_PERIODS))
    validation = selected.filter(F.col("periodo_archivo") == VALIDATION_PERIOD)
    return train, validation


def partition_checks(train, validation, checks: CheckLog) -> tuple[int, int]:
    from pyspark.sql import functions as F

    train_n = train.count()
    validation_n = validation.count()
    checks.require(train_n == TRAIN_ROWS, f"entrenamiento {train_n} != {TRAIN_ROWS}")
    checks.require(validation_n == VALIDATION_ROWS, f"validacion {validation_n} != {VALIDATION_ROWS}")
    checks.require(train.select("id_registro").distinct().count() == train_n, "entrenamiento: id_registro duplicado")
    checks.require(validation.select("id_registro").distinct().count() == validation_n, "validacion: id_registro duplicado")
    overlap = train.select("id_registro").intersect(validation.select("id_registro")).count()
    checks.require(overlap == 0, f"{overlap} id_registro compartidos entre entrenamiento y validacion")
    nulls = train.unionByName(validation).agg(
        *[F.sum(F.when(F.col(column).isNull(), 1).otherwise(0)).alias(column) for column in (TARGET, *PREDICTORS)]
    ).first().asDict()
    checks.require(all((value or 0) == 0 for value in nulls.values()), f"nulos en predictores u objetivo {nulls}")
    return train_n, validation_n


def categorical_stages() -> list[Any]:
    from pyspark.ml.feature import OneHotEncoder, StringIndexer

    indexers = [
        StringIndexer(inputCol=column, outputCol=f"{column}_idx", handleInvalid="keep", stringOrderType="alphabetAsc")
        for column in CATEGORICAL_PREDICTORS
    ]
    encoder = OneHotEncoder(
        inputCols=[f"{column}_idx" for column in CATEGORICAL_PREDICTORS],
        outputCols=[f"{column}_ohe" for column in CATEGORICAL_PREDICTORS],
        handleInvalid="keep",
        dropLast=True,
    )
    return [*indexers, encoder]


def assembled_inputs() -> list[str]:
    return [*NUMERIC_PREDICTORS, *[f"{column}_ohe" for column in CATEGORICAL_PREDICTORS]]


def linear_pipeline(config: dict[str, Any]):
    from pyspark.ml import Pipeline
    from pyspark.ml.feature import StandardScaler, VectorAssembler
    from pyspark.ml.regression import LinearRegression

    return Pipeline(
        stages=[
            *categorical_stages(),
            VectorAssembler(inputCols=assembled_inputs(), outputCol="features_raw"),
            StandardScaler(inputCol="features_raw", outputCol="features", withMean=True, withStd=True),
            LinearRegression(
                featuresCol="features",
                labelCol=TARGET,
                predictionCol="prediccion",
                regParam=config["regParam"],
                elasticNetParam=config["elasticNetParam"],
                standardization=False,
                fitIntercept=True,
                maxIter=LR_MAX_ITER,
                solver="l-bfgs",
            ),
        ]
    )


def forest_pipeline(config: dict[str, Any]):
    from pyspark.ml import Pipeline
    from pyspark.ml.feature import VectorAssembler
    from pyspark.ml.regression import RandomForestRegressor

    return Pipeline(
        stages=[
            *categorical_stages(),
            VectorAssembler(inputCols=assembled_inputs(), outputCol="features"),
            RandomForestRegressor(
                featuresCol="features",
                labelCol=TARGET,
                predictionCol="prediccion",
                numTrees=config["numTrees"],
                maxDepth=config["maxDepth"],
                maxBins=RF_MAX_BINS,
                subsamplingRate=1.0,
                featureSubsetStrategy="onethird",
                seed=SEED,
            ),
        ]
    )


def regression_metrics(frame, prediction_column: str) -> dict[str, float]:
    from pyspark.ml.evaluation import RegressionEvaluator

    return {
        metric: float(
            RegressionEvaluator(labelCol=TARGET, predictionCol=prediction_column, metricName=metric).evaluate(frame)
        )
        for metric in ("mae", "rmse", "r2")
    }


def readable_feature(name: str) -> str:
    for column in CATEGORICAL_PREDICTORS:
        prefix = f"{column}_ohe_"
        if name.startswith(prefix):
            code = name[len(prefix):]
            label = "categoria_no_vista" if code == "__unknown" else category_label(column, code)
            return f"{column}={label}"
    return name


def pipeline_feature_names(model, frame, vector_column: str) -> list[str]:
    metadata = model.transform(frame.limit(1)).schema[vector_column].metadata["ml_attr"]
    attributes = sorted((item for group in metadata["attrs"].values() for item in group), key=lambda item: item["idx"])
    return [readable_feature(item["name"]) for item in attributes]


def model_details(algorithm: str, model, frame) -> dict[str, Any]:
    names = pipeline_feature_names(model, frame, "features_raw" if algorithm == "regresion_lineal" else "features")
    estimator = model.stages[-1]
    if algorithm == "regresion_lineal":
        coefficients = estimator.coefficients.toArray().tolist()
        mapped = dict(zip(names, coefficients)) if len(names) == len(coefficients) else {}
        return {
            "intercepto": float(estimator.intercept),
            "coeficientes_espacio_estandarizado": mapped,
            "n_coeficientes": len(coefficients),
            "iteraciones": int(estimator.summary.totalIterations),
        }
    importances = estimator.featureImportances.toArray().tolist()
    mapped = dict(zip(names, importances)) if len(names) == len(importances) else {}
    return {
        "importancias": dict(sorted(mapped.items(), key=lambda item: -item[1])),
        "n_arboles": int(estimator.getNumTrees),
        "nodos_totales": int(sum(tree.numNodes for tree in estimator.trees)),
    }


def fit_candidates(algorithm: str, configs: list[dict[str, Any]], builder, train, validation) -> tuple[list[dict[str, Any]], Any]:
    rows = []
    best_model = None
    best_rmse = math.inf
    for config in configs:
        started = time.perf_counter()
        model = builder(config).fit(train)
        metrics = regression_metrics(model.transform(validation), "prediccion")
        row = {
            "algoritmo": algorithm,
            "config_id": config["config_id"],
            "parametros": {key: value for key, value in config.items() if key != "config_id"},
            **metrics,
            "n_validacion": VALIDATION_ROWS,
            "segundos": round(time.perf_counter() - started, 1),
            "seleccionado": False,
        }
        rows.append(row)
        print(f"  {config['config_id']}: MAE={metrics['mae']:,.2f} RMSE={metrics['rmse']:,.2f} R2={metrics['r2']:.4f}")
        if metrics["rmse"] < best_rmse:
            best_rmse = metrics["rmse"]
            best_model = model
    selected = min(rows, key=lambda item: item["rmse"])
    selected["seleccionado"] = True
    return rows, best_model


def figure_validation_rmse(rows: list[dict[str, Any]], baseline: dict[str, Any], path: Path) -> str:
    import matplotlib.pyplot as plt

    colors = {"regresion_lineal": "#4C72B0", "random_forest": "#55A868"}
    figure, axis = plt.subplots(figsize=(10, 4.8))
    labels = [row["config_id"] for row in rows]
    bars = axis.bar(labels, [row["rmse"] for row in rows], color=[colors[row["algoritmo"]] for row in rows])
    for bar, row in zip(bars, rows):
        if row["seleccionado"]:
            bar.set_edgecolor("black")
            bar.set_linewidth(2)
    axis.axhline(baseline["rmse"], color="#C44E52", linestyle="--", label=f"baseline media T1-T3 (RMSE {baseline['rmse']:,.0f})")
    axis.set_ylabel("RMSE en 2025T4 (Q)")
    axis.set_title("Validación temporal 2025T4: RMSE por configuración (borde negro = seleccionada)")
    axis.tick_params(axis="x", rotation=30)
    axis.legend()
    return save_figure(figure, path)


def run_models_validation() -> int:
    from pyspark.sql import functions as F

    started = time.perf_counter()
    checks = CheckLog()
    validation_path = FOUNDATION_OUT / "validation_foundation.json"
    if not (validation_path.exists() and read_json(validation_path).get("ok") is True and parquet_complete(ELIGIBLE_2025)):
        print("MODELOS_VALIDACION=FALLA la fundacion no esta validada: ejecute --validate-foundation")
        return 1
    MODELS_VALIDATION_OUT.mkdir(parents=True, exist_ok=True)
    spark = build_spark("lab7-models-validation")
    spark.sparkContext.setLogLevel("WARN")
    train = validation = None
    try:
        train, validation = temporal_partitions(spark, checks)
        train = train.cache()
        validation = validation.cache()
        train_n, validation_n = partition_checks(train, validation, checks)
        if not checks.ok:
            print("MODELOS_VALIDACION=FALLA particiones temporales invalidas")
            return 1
        train_mean = float(train.agg(F.avg(TARGET)).first()[0])
        baseline_frame = validation.withColumn("pred_baseline", F.lit(train_mean))
        baseline = {
            "algoritmo": "baseline_media",
            "config_id": "baseline_media_T1_T3",
            "parametros": {"media_entrenamiento": train_mean},
            **regression_metrics(baseline_frame, "pred_baseline"),
            "n_validacion": validation_n,
            "seleccionado": True,
        }
        print(f"Baseline media T1-T3 = Q{train_mean:,.2f}: MAE={baseline['mae']:,.2f} RMSE={baseline['rmse']:,.2f} R2={baseline['r2']:.4f}")
        print("Regresion lineal (StandardScaler explicito, standardization=False)")
        lr_rows, lr_model = fit_candidates("regresion_lineal", LR_CONFIGS, linear_pipeline, train, validation)
        print(f"Random Forest (semilla {SEED}, sin estandarizacion)")
        rf_rows, rf_model = fit_candidates("random_forest", RF_CONFIGS, forest_pipeline, train, validation)
        lr_selected = next(row for row in lr_rows if row["seleccionado"])
        rf_selected = next(row for row in rf_rows if row["seleccionado"])
        lr_model.write().overwrite().save(str(LR_VALIDATION_MODEL))
        rf_model.write().overwrite().save(str(RF_VALIDATION_MODEL))
        lr_predictions = lr_model.transform(validation).select("id_registro", F.col("prediccion").alias("pred_regresion_lineal"))
        rf_predictions = rf_model.transform(validation).select("id_registro", F.col("prediccion").alias("pred_random_forest"))
        predictions = (
            baseline_frame.select("id_registro", "periodo_archivo", "fila_origen", TARGET, "pred_baseline")
            .join(lr_predictions, "id_registro", "inner")
            .join(rf_predictions, "id_registro", "inner")
            .withColumn("residuo_baseline", F.col(TARGET) - F.col("pred_baseline"))
            .withColumn("residuo_regresion_lineal", F.col(TARGET) - F.col("pred_regresion_lineal"))
            .withColumn("residuo_random_forest", F.col(TARGET) - F.col("pred_random_forest"))
            .orderBy("fila_origen")
        )
        predictions.coalesce(1).write.mode("overwrite").parquet(str(VALIDATION_PREDICTIONS))
        persisted = spark.read.parquet(str(VALIDATION_PREDICTIONS))
        checks.require(persisted.count() == validation_n, "predicciones persistidas con cobertura incompleta")
        candidates = [baseline, lr_selected, rf_selected]
        leader = min(candidates, key=lambda item: item["rmse"])
        figure = figure_validation_rmse(lr_rows + rf_rows, baseline, MODELS_VALIDATION_OUT / "figures" / "rmse_validacion_2025T4.png")
        write_json(MODELS_VALIDATION_OUT / "validation_configs.json", {
            "particion": {
                "entrenamiento": list(TRAIN_PERIODS),
                "validacion": VALIDATION_PERIOD,
                "n_entrenamiento": train_n,
                "n_validacion": validation_n,
                "datos_2026_usados": False,
                "randomSplit": False,
                "cross_validation": False,
            },
            "objetivo": TARGET,
            "predictores_numericos": list(NUMERIC_PREDICTORS),
            "predictores_categoricos": list(CATEGORICAL_PREDICTORS),
            "preprocesamiento_categorico": "StringIndexer(handleInvalid=keep, stringOrderType=alphabetAsc) + OneHotEncoder(handleInvalid=keep, dropLast=True)",
            "regresion_lineal": {
                "estandarizacion": "StandardScaler(withMean=True, withStd=True) sobre el vector ensamblado; LinearRegression(standardization=False)",
                "solver": "l-bfgs",
                "max_iter": LR_MAX_ITER,
                "configuraciones": LR_CONFIGS,
                "escala_regParam": "Spark expresa regParam en unidades del objetivo (penalizacion efectiva regParam/desviacion del salario)",
                "codificacion": "OneHotEncoder con handleInvalid=keep conserva todas las categorias vistas y una columna __unknown; la colinealidad con el intercepto se controla con ridge",
            },
            "random_forest": {
                "estandarizacion": "ninguna",
                "semilla": SEED,
                "max_bins": RF_MAX_BINS,
                "feature_subset_strategy": "onethird",
                "configuraciones": RF_CONFIGS,
            },
            "criterio_seleccion": "menor RMSE en 2025T4 por algoritmo",
            "ajuste_preprocesamiento": "solo 2025T1-T3 dentro de cada Pipeline",
            "recorte_objetivo_o_predicciones": False,
        })
        write_json(MODELS_VALIDATION_OUT / "validation_metrics.json", {"baseline": baseline, "candidatos": lr_rows + rf_rows})
        summary = {
            "generado_utc": utc_now(),
            "segundos": round(time.perf_counter() - started, 1),
            "spark": spark.version,
            "codigo_sha256": sha256_file(Path(__file__)),
            "entrada": {"ruta": relative(ELIGIBLE_2025), **parquet_fingerprint(ELIGIBLE_2025)},
            "n_entrenamiento": train_n,
            "n_validacion": validation_n,
            "media_entrenamiento": train_mean,
            "seleccion": {
                "baseline_media": baseline,
                "regresion_lineal": lr_selected,
                "random_forest": rf_selected,
            },
            "lider_provisional_2025T4": {"algoritmo": leader["algoritmo"], "config_id": leader["config_id"], "rmse": leader["rmse"]},
            "detalles_modelos": {
                "regresion_lineal": model_details("regresion_lineal", lr_model, validation),
                "random_forest": model_details("random_forest", rf_model, validation),
            },
            "salidas": {
                "predicciones": relative(VALIDATION_PREDICTIONS),
                "modelo_regresion_lineal": relative(LR_VALIDATION_MODEL),
                "modelo_random_forest": relative(RF_VALIDATION_MODEL),
                "figura_rmse": figure,
            },
            "verificacion": checks.as_dict(),
        }
        write_json(MODELS_VALIDATION_OUT / "models_validation_summary.json", summary)
        print(f"Seleccion LR={lr_selected['config_id']} RMSE={lr_selected['rmse']:,.2f} | RF={rf_selected['config_id']} RMSE={rf_selected['rmse']:,.2f}")
        print(f"Lider provisional 2025T4: {leader['config_id']} (RMSE {leader['rmse']:,.2f})")
        status = "OK" if checks.ok else "FALLA"
        print(f"MODELOS_VALIDACION={status} controles={checks.passed} fallas={len(checks.failures)}")
        return 0 if checks.ok else 1
    finally:
        for frame in (train, validation):
            if frame is not None:
                frame.unpersist()
        spark.catalog.clearCache()
        spark.stop()


def validate_validation_metadata(configs: dict[str, Any], metrics: dict[str, Any], summary: dict[str, Any], checks: CheckLog) -> None:
    partition = configs["particion"]
    checks.require(partition["entrenamiento"] == list(TRAIN_PERIODS) and partition["validacion"] == VALIDATION_PERIOD, "particion registrada")
    checks.require(not partition["datos_2026_usados"] and not partition["randomSplit"] and not partition["cross_validation"], "particion no temporal")
    checks.require(summary["n_entrenamiento"] == TRAIN_ROWS and summary["n_validacion"] == VALIDATION_ROWS, "conteos del resumen")
    checks.require(configs["predictores_numericos"] == list(NUMERIC_PREDICTORS), "predictores numericos")
    checks.require(configs["predictores_categoricos"] == list(CATEGORICAL_PREDICTORS), "predictores categoricos")
    checks.require(configs["objetivo"] == TARGET, "objetivo")
    candidates = metrics["candidatos"]
    lr_rows = [row for row in candidates if row["algoritmo"] == "regresion_lineal"]
    rf_rows = [row for row in candidates if row["algoritmo"] == "random_forest"]
    checks.require(len({row["parametros"]["regParam"] for row in lr_rows}) >= 3, "menos de tres regularizaciones lineales")
    checks.require(all(0 <= row["parametros"]["regParam"] <= max(LR_REG_PARAMS) for row in lr_rows), "regularizacion no acotada")
    checks.require(len({(row["parametros"]["numTrees"], row["parametros"]["maxDepth"]) for row in rf_rows}) >= 2, "menos de dos configuraciones RF")
    checks.require(configs["random_forest"]["semilla"] == SEED, "semilla RF")
    checks.require("standardization=False" in configs["regresion_lineal"]["estandarizacion"], "doble estandarizacion lineal")
    for row in candidates:
        checks.require(row["n_validacion"] == VALIDATION_ROWS, f"{row['config_id']}: n de validacion")
        checks.require(row["rmse"] >= row["mae"] >= 0, f"{row['config_id']}: MAE/RMSE incoherentes")
    for algorithm, rows in (("regresion_lineal", lr_rows), ("random_forest", rf_rows)):
        selected = [row for row in rows if row["seleccionado"]]
        checks.require(len(selected) == 1, f"{algorithm}: seleccion no unica")
        checks.require(selected[0]["rmse"] == min(row["rmse"] for row in rows), f"{algorithm}: seleccion no es el menor RMSE")
        checks.require(summary["seleccion"][algorithm]["config_id"] == selected[0]["config_id"], f"{algorithm}: resumen distinto")
    selections = summary["seleccion"]
    leader = min(selections.values(), key=lambda item: item["rmse"])
    checks.require(summary["lider_provisional_2025T4"]["config_id"] == leader["config_id"], "lider provisional incorrecto")
    checks.require(
        close_to(metrics["baseline"]["parametros"]["media_entrenamiento"], summary["media_entrenamiento"]), "media del baseline"
    )


def run_validate_models_validation() -> int:
    from pyspark.ml import PipelineModel
    from pyspark.sql import functions as F

    checks = CheckLog()
    names = ("validation_configs.json", "validation_metrics.json", "models_validation_summary.json")
    absent = [relative(MODELS_VALIDATION_OUT / name) for name in names if not (MODELS_VALIDATION_OUT / name).exists()]
    absent += [relative(VALIDATION_PREDICTIONS)] if not parquet_complete(VALIDATION_PREDICTIONS) else []
    absent += [relative(path) for path in (LR_VALIDATION_MODEL, RF_VALIDATION_MODEL) if not (path / "metadata").exists()]
    if absent:
        print(f"VALIDACION_MODELOS=FALLA artefactos ausentes: {absent}")
        return 1
    configs, metrics, summary = (read_json(MODELS_VALIDATION_OUT / name) for name in names)
    print("Metadatos de seleccion:")
    validate_validation_metadata(configs, metrics, summary, checks)
    checks.require(summary["verificacion"]["ok"], "el resumen registra fallas")
    checks.require(summary["entrada"]["sha256"] == parquet_fingerprint(ELIGIBLE_2025)["sha256"], "eligible_2025 cambio despues de la seleccion")
    figure = ROOT / summary["salidas"]["figura_rmse"]
    checks.require(figure.exists() and figure.read_bytes()[:8] == b"\x89PNG\r\n\x1a\n", "figura RMSE ausente")
    spark = build_spark("lab7-validate-models-validation")
    spark.sparkContext.setLogLevel("WARN")
    try:
        train, validation = temporal_partitions(spark, checks)
        partition_checks(train, validation, checks)
        train_mean = float(train.agg(F.avg(TARGET)).first()[0])
        checks.require(close_to(train_mean, summary["media_entrenamiento"], 1e-9), "media T1-T3 recalculada distinta")
        predictions = spark.read.parquet(str(VALIDATION_PREDICTIONS)).cache()
        total = predictions.count()
        checks.require(total == VALIDATION_ROWS, f"predicciones {total} != {VALIDATION_ROWS}")
        checks.require(predictions.select("id_registro").distinct().count() == total, "predicciones: id_registro duplicado")
        checks.require(predictions.filter(F.col("periodo_archivo") != VALIDATION_PERIOD).count() == 0, "predicciones fuera de 2025T4")
        checks.require(predictions.filter(F.col("id_registro").startswith("2026")).count() == 0, "predicciones con registros 2026")
        missing = validation.select("id_registro").subtract(predictions.select("id_registro")).count()
        extra = predictions.select("id_registro").subtract(validation.select("id_registro")).count()
        checks.require(missing == 0 and extra == 0, f"cobertura distinta de T4: faltan {missing}, sobran {extra}")
        prediction_columns = ("pred_baseline", "pred_regresion_lineal", "pred_random_forest")
        nulls = predictions.agg(*[F.sum(F.when(F.col(c).isNull() | F.isnan(c), 1).otherwise(0)).alias(c) for c in prediction_columns]).first().asDict()
        checks.require(all((value or 0) == 0 for value in nulls.values()), f"predicciones nulas {nulls}")
        baseline_values = predictions.agg(F.min("pred_baseline"), F.max("pred_baseline")).first()
        checks.require(
            close_to(baseline_values[0], train_mean, 1e-9) and close_to(baseline_values[1], train_mean, 1e-9),
            "baseline distinto de la media de entrenamiento",
        )
        target_mismatch = (
            predictions.join(validation.select("id_registro", F.col(TARGET).alias("objetivo_original")), "id_registro")
            .filter(F.col(TARGET) != F.col("objetivo_original"))
            .count()
        )
        checks.require(target_mismatch == 0, f"{target_mismatch} objetivos alterados o recortados")
        for algorithm, column in (("baseline_media", "pred_baseline"), ("regresion_lineal", "pred_regresion_lineal"), ("random_forest", "pred_random_forest")):
            recomputed = regression_metrics(predictions, column)
            stored = summary["seleccion"][algorithm]
            for metric, value in recomputed.items():
                checks.require(close_to(value, stored[metric], 1e-6), f"{algorithm}: {metric} recalculado {value} != {stored[metric]}")
            print(f"  {algorithm}: MAE={recomputed['mae']:,.2f} RMSE={recomputed['rmse']:,.2f} R2={recomputed['r2']:.4f}")
        for column, path in (("pred_regresion_lineal", LR_VALIDATION_MODEL), ("pred_random_forest", RF_VALIDATION_MODEL)):
            reloaded = PipelineModel.load(str(path)).transform(validation).select("id_registro", F.col("prediccion").alias("recargada"))
            difference = predictions.join(reloaded, "id_registro").agg(F.max(F.abs(F.col(column) - F.col("recargada")))).first()[0]
            checks.require(difference is not None and difference <= 1e-6, f"{relative(path)}: predicciones recargadas difieren {difference}")
        predictions.unpersist()
    finally:
        spark.stop()
    write_json(MODELS_VALIDATION_OUT / "validation_models_validation.json", {"generado_utc": utc_now(), **checks.as_dict()})
    status = "OK" if checks.ok else "FALLA"
    print(f"Lider provisional 2025T4: {summary['lider_provisional_2025T4']['config_id']}")
    print(f"VALIDACION_MODELOS={status} controles={checks.passed} fallas={len(checks.failures)} advertencias={len(checks.warnings)}")
    return 0 if checks.ok else 1


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Laboratorio 7 - flujo final reproducible con Spark")
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--preflight", action="store_true", help="verifica entorno, fuentes, diccionarios y Parquet")
    group.add_argument("--foundation", action="store_true", help="carga selectiva, auditoria, filtros y Parquet")
    group.add_argument("--validate-foundation", action="store_true", help="valida solo artefactos persistidos")
    group.add_argument("--eda-clustering", action="store_true", help="EDA, correlacion y KMeans con elegibles 2025")
    group.add_argument("--validate-eda-clustering", action="store_true", help="valida artefactos persistidos de EDA y KMeans")
    group.add_argument("--models-validation", action="store_true", help="baseline, regresion lineal y Random Forest: T1-T3 vs 2025T4")
    group.add_argument("--validate-models-validation", action="store_true", help="valida seleccion temporal persistida")
    parser.add_argument("--force", action="store_true", help="reprocesa periodos aunque existan artefactos vigentes")
    parser.add_argument("--only", nargs="+", choices=list(SOURCES_BY_PERIOD), help="procesa solo estos periodos")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    if args.preflight:
        return run_preflight()
    if args.foundation:
        return run_foundation(force=args.force, only=args.only)
    if args.eda_clustering:
        return run_eda_clustering()
    if args.validate_eda_clustering:
        return run_validate_eda_clustering()
    if args.models_validation:
        return run_models_validation()
    if args.validate_models_validation:
        return run_validate_models_validation()
    return run_validate_foundation()


if __name__ == "__main__":
    sys.exit(main())
